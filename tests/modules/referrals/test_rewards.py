# ruff: noqa: E501
"""Epic 42 story 42.4 (FR-216, AD-23, AD-24): the acceptance-criteria flow -- a ₹25,000 senior program with a 90-day
Active condition: a Worker refers by resume upload, the friend accepts (application ``source='referral'``), the
referrer sees In process then Hired; after 90 Active days the reward is Eligible, a Finance User approves it (*),
and it appears as a one-off line in the next Payroll Export (adjustments column, fixed amount, no
``missing_pay_rate``); finalizing marks it paid. Also: condition not met (ended day 60 / dropped) -> void, human-only
approval, payout export (formula escaping) + mark paid, manual void, erasure."""

import csv
import io
import uuid
from datetime import date, timedelta
from typing import Any

import pytest
from fastapi import FastAPI

from app.core.actors import ActorRef
from app.core.database import tenant_session
from app.core.errors import Forbidden
from app.models.referrals import ReferralReward
from app.services.privacy.registry import ErasureContext
from app.services.referrals import rewards
from app.tasks.referrals import _placement, evaluate_tenant
from tests.modules.demands.helpers import DOrg, h
from tests.modules.identity.conftest import join
from tests.modules.payroll.test_payroll import P, add_bank, finalize, le_of, month_start, preview
from tests.modules.placements.helpers import PLACEMENTS, direct_placement, waive_all
from tests.modules.referrals.helpers import (
    API,
    ME,
    WORK,
    accept,
    create_program,
    employee,
    execute,
    refer,
    refer_ok,
    referee_token,
    rows,
    scalar,
    senior_demand,
)
from tests.p2 import sign_in_candidate

TODAY = date.today()


async def _hire(org: DOrg, referral_id: str, demand: dict[str, Any], *, join_today: bool = True) -> dict[str, Any]:
    cand, app_id = (await rows(org, "SELECT candidate_id, application_id FROM referrals WHERE id = :i", i=referral_id))[0]
    p = await direct_placement(org, cand, application_id=str(app_id), demand_id=demand["id"])
    if join_today:
        await waive_all(org, p["id"])
        r = await org.admin.client.post(
            f"{PLACEMENTS}/{p['id']}/join", json={"actual_start_date": TODAY.isoformat()}, headers=h(org.admin)
        )
        assert r.status_code == 200, r.text
    return p


async def _user_referral_hired(app: FastAPI, org: DOrg) -> tuple[Any, dict[str, Any], dict[str, Any]]:
    await create_program(org)
    emp = await employee(app, org)
    demand = await senior_demand(org)
    email = f"hire-{uuid.uuid4().hex[:6]}@friends.example.com"
    out = await refer_ok(emp, demand["id"], email=email)
    assert (await accept(emp.client, referee_token(email))).status_code == 200
    p = await _hire(org, out["id"], demand)
    return emp, out, p


async def _reward(org: DOrg, referral_id: str) -> Any:
    return (
        await rows(
            org,
            "SELECT status, amount, condition_due_on, payout_channel, approval_request_id, void_reason, id "
            "FROM referral_rewards WHERE referral_id = :i",
            i=referral_id,
        )
    )[0]


async def _finance(app: FastAPI, org: DOrg) -> Any:
    return await join(app, org.tenant, "finance", name="Fiona Finance")


async def test_ac_worker_referral_to_payroll(app: FastAPI, dorg: DOrg) -> None:
    le_id = await le_of(dorg)
    cal = await dorg.admin.client.put(
        f"{P}/calendars/{le_id}",
        json={"cutoff_day": "last", "generate_from": (month_start() - timedelta(days=40)).isoformat()},
        headers=h(dorg.admin),
    )
    assert cal.status_code == 200, cal.text
    await create_program(dorg, eligible_referrers="both")
    demand = await senior_demand(dorg)
    worker_p = await direct_placement(dorg, planned_start_date=(month_start() - timedelta(days=35)).isoformat())
    wid = uuid.UUID(worker_p["candidate_id"])
    await add_bank(dorg, wid)
    worker = await sign_in_candidate(app, wid, dorg.tenant_id)
    try:
        email = f"ac-{uuid.uuid4().hex[:6]}@friends.example.com"
        made = await refer(worker, demand["id"], url=WORK, email=email)
        assert made.status_code == 201, made.text
        rid = made.json()["id"]
        assert (await accept(worker, referee_token(email))).status_code == 200
        assert await scalar(dorg, "SELECT source FROM applications WHERE source_ref = :r", r=rid) == "referral"
        assert (await worker.get(WORK)).json()[0]["status_label"] == "In process"

        await _hire(dorg, rid, demand)
        mine = (await worker.get(WORK)).json()[0]
        assert mine["status_label"] == "Hired"
        assert mine["reward"] == {"status_label": "Pending", "amount": "25000.00", "currency": "INR"}
        st, amount, due, channel, req, _, reward_id = await _reward(dorg, rid)
        assert (st, str(amount), due, channel, req) == ("pending_condition", "25000.00", TODAY + timedelta(days=90), "payroll", None)

        assert await evaluate_tenant(dorg.tenant_id, TODAY + timedelta(days=89)) == 0
        assert await evaluate_tenant(dorg.tenant_id, TODAY + timedelta(days=90)) == 1
        st, _, _, _, req, _, _ = await _reward(dorg, rid)
        assert st == "eligible" and req is not None
        assert (await worker.get(WORK)).json()[0]["reward"]["status_label"] == "Eligible"

        fin = await _finance(app, dorg)
        dec = await fin.client.post(f"{API}/approvals/{req}/decide", json={"action": "approve"}, headers=fin.h())
        assert dec.status_code == 200, dec.text
        assert (await _reward(dorg, rid))[0] == "approved"

        period = await dorg.admin.client.get(f"{P}/periods", params={"legal_entity_id": str(le_id)})
        period_id = next(r["id"] for r in period.json() if r["period_start"] == month_start().isoformat())
        pv = await preview(_W(dorg), period_id)
        assert pv["totals"]["INR"]["gross_earnings_inputs"] == "25000.00", pv
        assert not any(pv.get("exceptions", {}).get("missing_pay_rate", []) or [])
        prow = (await dorg.admin.client.get(f"{P}/versions/{pv['id']}/rows")).json()
        assert len(prow) == 1 and prow[0]["adjustments_total"] == "25000.00" and prow[0]["candidate_id"] == str(wid)
        fin_resp = await finalize(_W(dorg), pv["id"])
        assert fin_resp.status_code == 200, fin_resp.text
        st, *_rest = await _reward(dorg, rid)
        assert st == "paid"
        assert await scalar(dorg, "SELECT payout_ref FROM referral_rewards WHERE id = :i", i=reward_id) == fin_resp.json()["id"]
        assert (await worker.get(WORK)).json()[0]["reward"]["status_label"] == "Paid"
    finally:
        await worker.aclose()


class _W:  # noqa: B903
    """Minimal stand-in for the payroll test ``World`` (``preview`` / ``finalize`` only read ``org``)."""

    def __init__(self, org: DOrg) -> None:
        self.org = org


async def test_ended_before_condition_voids_and_dropped_voids(app: FastAPI, dorg: DOrg) -> None:
    _, out, p = await _user_referral_hired(app, dorg)
    assert (await _reward(dorg, out["id"]))[3] == "payout_export"  # a User referrer
    await execute(dorg, "UPDATE placements SET status = 'ended', end_date = :d WHERE id = :i", d=TODAY + timedelta(days=60), i=p["id"])
    await _placement(dorg.tenant_id, uuid.UUID(p["id"]), "ended", (TODAY + timedelta(days=60)).isoformat())
    st, *_, void_reason, _ = await _reward(dorg, out["id"])
    assert (st, void_reason) == ("void", "condition_not_met")
    assert await evaluate_tenant(dorg.tenant_id, TODAY + timedelta(days=120)) == 0

    _, out2, p2 = await _user_referral_hired(app, dorg)
    drop = await dorg.admin.client.post(f"{PLACEMENTS}/{p2['id']}/drop", json={"reason": "no_show"}, headers=h(dorg.admin))
    if drop.status_code != 200:  # an active placement cannot be dropped through the API: drive the handler
        await execute(dorg, "UPDATE placements SET status = 'dropped' WHERE id = :i", i=p2["id"])
        await _placement(dorg.tenant_id, uuid.UUID(p2["id"]), "dropped", None)
    st, *_, void_reason, _ = await _reward(dorg, out2["id"])
    assert (st, void_reason) == ("void", "placement_dropped")
    assert await scalar(dorg, "SELECT status FROM referrals WHERE id = :i", i=out2["id"]) == "not_selected"


async def test_duplicate_or_no_program_creates_no_reward(app: FastAPI, dorg: DOrg) -> None:
    emp = await employee(app, dorg)
    demand = await senior_demand(dorg)  # no program at all
    email = f"np-{uuid.uuid4().hex[:6]}@friends.example.com"
    out = await refer_ok(emp, demand["id"], email=email)
    assert (await accept(emp.client, referee_token(email))).status_code == 200
    await _hire(dorg, out["id"], demand, join_today=False)
    assert await scalar(dorg, "SELECT status FROM referrals WHERE id = :i", i=out["id"]) == "hired"
    assert await scalar(dorg, "SELECT count(*) FROM referral_rewards WHERE referral_id = :i", i=out["id"]) == 0
    assert (await emp.client.get(ME)).json()[0]["reward"] is None


@pytest.mark.parametrize("kind", ["ai", "workflow", "system", "api_key"])
async def test_reward_approval_is_human_only(app: FastAPI, dorg: DOrg, kind: str) -> None:
    from app.schemas.approvals import DecideIn
    from app.services import approvals
    from tests.modules.demands.helpers import admin_ctx

    _, out, _ = await _user_referral_hired(app, dorg)
    assert await evaluate_tenant(dorg.tenant_id, TODAY + timedelta(days=90)) == 1
    st, _, _, _, req, _, reward_id = await _reward(dorg, out["id"])
    assert st == "eligible"
    async with tenant_session(dorg.tenant_id) as db:
        with pytest.raises(Forbidden):
            await approvals.decide(admin_ctx(dorg), db, req, DecideIn(action="approve"), actor=ActorRef(kind, None))  # type: ignore[arg-type]
    async with tenant_session(dorg.tenant_id) as db:
        reward = await db.get(ReferralReward, reward_id)
        assert reward is not None
        with pytest.raises(Forbidden):
            await rewards.mark_approved(db, reward, ActorRef(kind, None))  # type: ignore[arg-type]
    assert (await _reward(dorg, out["id"]))[0] == "eligible"


async def test_human_only_registered_and_rejection_voids(app: FastAPI, dorg: DOrg) -> None:
    from app.services.workflow.catalogue import is_human_only

    assert is_human_only("referral_reward", "approved")
    _, out, _ = await _user_referral_hired(app, dorg)
    await evaluate_tenant(dorg.tenant_id, TODAY + timedelta(days=90))
    req = (await _reward(dorg, out["id"]))[4]
    fin = await _finance(app, dorg)
    r = await fin.client.post(f"{API}/approvals/{req}/decide", json={"action": "reject", "comment": "Duplicate claim"}, headers=fin.h())
    assert r.status_code == 200, r.text
    st, *_, void_reason, _ = await _reward(dorg, out["id"])
    assert (st, void_reason) == ("void", "rejected")


async def test_payout_export_escapes_formulas_and_mark_paid(app: FastAPI, dorg: DOrg) -> None:
    emp, out, _ = await _user_referral_hired(app, dorg)
    await execute(dorg, "UPDATE referrals SET referee_name = '=HYPERLINK(\"http://x\")' WHERE id = :i", i=out["id"])
    await evaluate_tenant(dorg.tenant_id, TODAY + timedelta(days=90))
    req = (await _reward(dorg, out["id"]))[4]
    fin = await _finance(app, dorg)
    assert (await fin.client.post(f"{API}/approvals/{req}/decide", json={"action": "approve"}, headers=fin.h())).status_code == 200
    csv_resp = await dorg.admin.client.get(f"{API}/referral-rewards/payout-export", params={"format": "csv"})
    assert csv_resp.status_code == 200, csv_resp.text
    data = list(csv.reader(io.StringIO(csv_resp.text)))
    assert data[0][0] == "Reward ID" and len(data) == 2
    assert data[1][5].startswith("'=") and data[1][6] == "25000.00" and data[1][3] == "Esha Employee"
    xlsx = await dorg.admin.client.get(f"{API}/referral-rewards/payout-export", params={"format": "xlsx"})
    assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws is not None and ws.cell(row=2, column=6).data_type == "s"
    reward_id = (await _reward(dorg, out["id"]))[6]
    no_perm = await emp.client.get(f"{API}/referral-rewards/payout-export")
    assert no_perm.status_code == 403
    paid = await dorg.admin.client.post(
        f"{API}/referral-rewards/mark-paid",
        json={"ids": [str(reward_id)], "paid_on": TODAY.isoformat(), "payout_ref": "NEFT-2026-10-001"},
        headers=h(dorg.admin),
    )
    assert paid.status_code == 200 and paid.json()["updated"] == 1
    assert (await _reward(dorg, out["id"]))[0] == "paid"
    again = await dorg.admin.client.post(
        f"{API}/referral-rewards/mark-paid",
        json={"ids": [str(reward_id)], "paid_on": TODAY.isoformat(), "payout_ref": "x"},
        headers=h(dorg.admin),
    )
    assert again.status_code == 409
    mine = (await emp.client.get(ME)).json()[0]
    assert mine["reward"]["status_label"] == "Paid"


async def test_manual_void_and_list(app: FastAPI, dorg: DOrg) -> None:
    _, out, _ = await _user_referral_hired(app, dorg)
    reward_id = (await _reward(dorg, out["id"]))[6]
    listed = (await dorg.admin.client.get(f"{API}/referral-rewards", params={"status": "pending_condition"})).json()
    assert [r["id"] for r in listed] == [str(reward_id)] and listed[0]["amount"] == "25000.00"
    v = await dorg.admin.client.post(f"{API}/referral-rewards/{reward_id}/void", json={"reason": "Policy exception"}, headers=h(dorg.admin))
    assert v.status_code == 200 and v.json()["status"] == "void" and v.json()["void_reason"] == "manual"
    again = await dorg.admin.client.post(f"{API}/referral-rewards/{reward_id}/void", json={"reason": "x"}, headers=h(dorg.admin))
    assert again.status_code == 409


async def test_erasure_of_referee_keeps_reward(app: FastAPI, dorg: DOrg) -> None:
    from app.services.referrals_erasure import ReferralsHandler

    _, out, _ = await _user_referral_hired(app, dorg)
    cand = await scalar(dorg, "SELECT candidate_id FROM referrals WHERE id = :i", i=out["id"])
    handler = ReferralsHandler()
    ectx = ErasureContext(tenant_id=dorg.tenant_id, candidate_id=cand, run_id=uuid.uuid4(), reason="dsr")
    async with tenant_session(dorg.tenant_id) as db:
        exported = await handler.export(db, ectx)
        assert exported is not None and exported["referrals_received"][0]["status"] == "Hired"
        counts = await handler.erase(db, ectx)
        assert counts["referrals"] == 1
        assert await handler.verify(db, ectx) == []
    row = (await rows(dorg, "SELECT referee_name, referee_email, status FROM referrals WHERE id = :i", i=out["id"]))[0]
    assert row == ("[erased]", None, "hired")
    assert await scalar(dorg, "SELECT amount FROM referral_rewards WHERE referral_id = :i", i=out["id"]) is not None
