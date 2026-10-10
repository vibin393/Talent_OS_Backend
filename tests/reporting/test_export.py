"""Exports (A5, FR-61): CSV / XLSX round-trips, exported rows = listed rows, size limits, masking, formula
injection, async runs with requester-only downloads, expiry, revoked permissions, audit, 429, cross-tenant."""

import csv
import io
import uuid
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text, update

from app.core.database import tenant_session
from app.core.settings import get_settings
from app.models.identity import Role
from app.models.reporting import Export
from app.services.reporting import export as export_engine
from tests.reporting.helpers import API, ROrg, Seed, ist
from tests.reporting.test_lists import reporter
from tests.reporting.timeline import RANGE, build_march

EXPORTS = f"{API}/exports"


async def post_export(
    m: Any, kind: str, fmt: str, filters: dict[str, Any] | None = None, **extra: Any
) -> Any:
    return await m.client.post(
        EXPORTS, json={"kind": kind, "format": fmt, "filters": filters or {}, **extra}, headers=m.h()
    )


def parse_csv(content: bytes) -> list[list[str]]:
    assert content.startswith(b"\xef\xbb\xbf"), "UTF-8 BOM"
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))


async def audit_rows(tenant_id: uuid.UUID, action: str) -> list[Any]:
    async with tenant_session(tenant_id) as db:
        rows = await db.execute(
            text("SELECT entity_type, entity_id, data FROM audit_logs WHERE action = :a ORDER BY occurred_at"),
            {"a": action},
        )
        return list(rows.mappings().all())


async def test_sync_csv_matches_list_rows(rorg: ROrg) -> None:
    await build_march(rorg)
    resp = await post_export(rorg.admin, "submissions", "csv", RANGE)
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("text/csv")
    assert resp.headers["cache-control"] == "no-store"
    assert resp.headers["content-disposition"].startswith('attachment; filename="submissions-')
    assert resp.headers["x-tos-masked-columns"] == ""
    rows = parse_csv(resp.content)
    header, data = rows[0], rows[1:]
    assert header[:3] == ["Submission ID", "Candidate", "REQ"]
    assert "Bill rate" in header and "Bill rate currency" in header
    listed = (
        await rorg.admin.client.get(f"{API}/reports/lists/submissions", params={**RANGE, "page_size": 100})
    ).json()
    assert len(data) == listed["total"] == 5
    assert [r[0] for r in data] == [i["human_id"] for i in listed["items"]]
    sent = header.index("Sent")
    assert all(len(r[sent]) == 16 for r in data)  # yyyy-mm-dd HH:MM in the tenant timezone
    assert any(r[sent] == "2026-03-09 00:30" for r in data)
    bill = header.index("Bill rate")
    assert data[0][bill] == "1500.00" and data[0][bill + 1] == "INR"
    status = header.index("Status")
    assert {r[status] for r in data} == {"Sent"}  # enum label
    audits = await audit_rows(rorg.tenant_id, "export.run")
    assert len(audits) == 1
    assert audits[0]["entity_type"] == "submissions"
    assert audits[0]["data"]["mode"] == "sync" and audits[0]["data"]["row_count"] == 5
    async with tenant_session(rorg.tenant_id) as db:
        row = (await db.execute(select(Export))).scalar_one()
    assert (row.status, row.mode, row.file_id, row.row_count) == ("done", "sync", None, 5)


async def test_sync_xlsx_round_trip(rorg: ROrg) -> None:
    await build_march(rorg)
    resp = await post_export(
        rorg.admin, "applications", "xlsx", RANGE, columns=["human_id", "candidate", "stage", "created_at"]
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-disposition"].endswith('.xlsx"')
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["applications"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Application ID", "Candidate", "Stage", "Added")
    assert ws["A1"].font.b is True
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == f"A1:D{len(rows)}"
    listed = (
        await rorg.admin.client.get(f"{API}/reports/lists/applications", params={**RANGE, "page_size": 100})
    ).json()
    assert [r[0] for r in rows[1:]] == [i["human_id"] for i in listed["items"]]


async def test_candidates_and_demands_exports_equal_their_lists(rorg: ROrg) -> None:
    await build_march(rorg)
    for kind, path, filters in (
        ("candidates", "/candidates", {}),
        ("demands", "/demands", {"status": ["open"]}),
    ):
        resp = await post_export(rorg.admin, kind, "csv", filters)
        assert resp.status_code == 200, resp.text
        data = parse_csv(resp.content)[1:]
        listed = (await rorg.admin.client.get(f"{API}{path}", params={**filters, "page_size": 100})).json()
        assert len(data) == listed["total"], kind
        first_key = "human_id" if kind == "candidates" else "req_id"
        assert [r[0] for r in data] == [i[first_key] for i in listed["items"]], kind


async def test_masking_removes_compensation_columns(app: Any, rorg: ROrg) -> None:
    await build_march(rorg)
    remi = await reporter(app, rorg)
    try:
        resp = await post_export(remi, "submissions", "xlsx", RANGE)
        assert resp.status_code == 200, resp.text
        masked = resp.headers["x-tos-masked-columns"].split(",")
        assert set(masked) == {"pay_rate", "bill_rate", "margin_amount", "margin_pct"}
        header = next(load_workbook(io.BytesIO(resp.content))["submissions"].iter_rows(values_only=True))
        assert not {"Pay rate", "Bill rate", "Margin", "Margin %"} & set(header)
        csv_resp = await post_export(remi, "submissions", "csv", RANGE)
        assert b"1500" not in csv_resp.content and b"Bill rate" not in csv_resp.content
        audits = await audit_rows(rorg.tenant_id, "export.run")
        assert set(audits[-1]["data"]["masked_columns"]) == set(masked)
    finally:
        await remi.client.aclose()


@pytest.mark.parametrize(
    "name", ['=HYPERLINK("http://evil.example","x")', "+1+2", "-3", "@SUM(A1)", "\tTabbed"]
)
async def test_formula_injection_guard(rorg: ROrg, name: str) -> None:
    async with Seed(rorg.tenant_id) as s:
        await s.candidate(rorg.admin.user_id, datetime.now(UTC), name)
    resp = await post_export(rorg.admin, "candidates", "csv", columns=["full_name"])
    assert parse_csv(resp.content)[1] == ["'" + name]
    resp = await post_export(rorg.admin, "candidates", "xlsx", columns=["full_name"])
    cell = load_workbook(io.BytesIO(resp.content))["candidates"]["A2"]
    assert cell.value == name and cell.data_type == "s"


async def test_size_limits(rorg: ROrg, monkeypatch: pytest.MonkeyPatch) -> None:
    await build_march(rorg)
    real = export_engine.count_rows

    async def fake(n: int) -> Any:
        async def count(db: Any, query: Any) -> int:
            await real(db, query)
            return n

        return count

    monkeypatch.setattr(export_engine, "count_rows", await fake(50_001))
    resp = await post_export(rorg.admin, "submissions", "csv")
    assert resp.status_code == 422
    err = resp.json()["error"]
    assert err["code"] == "export_too_large" and err["details"] == {"row_count": 50_001, "max": 50_000}
    async with tenant_session(rorg.tenant_id) as db:
        assert (await db.execute(select(Export))).first() is None  # nothing recorded
    monkeypatch.setattr(export_engine, "count_rows", await fake(5_000))
    assert (await post_export(rorg.admin, "submissions", "csv")).status_code == 200
    monkeypatch.setattr(export_engine, "count_rows", await fake(5_001))
    resp = await post_export(rorg.admin, "submissions", "csv")
    assert resp.status_code == 202, resp.text
    assert resp.json()["status"] == "queued"


async def _async_export(org: ROrg, m: Any, monkeypatch: pytest.MonkeyPatch, **filters: Any) -> str:
    monkeypatch.setattr(get_settings(), "export_sync_max_rows", 1)
    resp = await post_export(m, "submissions", "xlsx", filters or RANGE)
    assert resp.status_code == 202, resp.text
    return str(resp.json()["export_id"])


async def test_async_export_download_by_requester_only(rorg: ROrg, monkeypatch: pytest.MonkeyPatch) -> None:
    await build_march(rorg)
    export_id = await _async_export(rorg, rorg.admin, monkeypatch)
    out = (await rorg.admin.client.get(f"{EXPORTS}/{export_id}")).json()
    assert out["status"] == "done", out
    assert out["mode"] == "async" and out["row_count"] == 5
    assert out["download_url"] == f"/api/v1/exports/{export_id}/download"
    listing = (await rorg.admin.client.get(EXPORTS)).json()
    assert listing["items"][0]["id"] == export_id
    # The requester is notified in-app (Epic 10 ``export_ready``).
    async with tenant_session(rorg.tenant_id) as db:
        note = (
            (
                await db.execute(
                    text(
                        "SELECT type, link_path FROM notifications WHERE user_id = :u AND type = 'export_ready'"
                    ),
                    {"u": rorg.admin.user_id},
                )
            )
            .mappings()
            .all()
        )
    assert len(note) == 1 and note[0]["link_path"] == f"/exports?export={export_id}"
    # Download: 303 to the permission-checked file endpoint.
    resp = await rorg.admin.client.get(f"{EXPORTS}/{export_id}/download")
    assert resp.status_code == 303
    file_resp = await rorg.admin.client.get(resp.headers["location"])
    assert file_resp.status_code == 200
    rows = list(load_workbook(io.BytesIO(file_resp.content))["submissions"].iter_rows(values_only=True))
    assert len(rows) == 6
    # Another user (who may export too) gets 404 on both URLs.
    assert (await rorg.manager.client.get(f"{EXPORTS}/{export_id}")).status_code == 404
    assert (await rorg.manager.client.get(f"{EXPORTS}/{export_id}/download")).status_code == 404
    assert (await rorg.manager.client.get(resp.headers["location"])).status_code == 404
    downloads = await audit_rows(rorg.tenant_id, "export.downloaded")
    assert len(downloads) == 1
    async with tenant_session(rorg.tenant_id) as db:
        row = await db.get(Export, uuid.UUID(export_id))
        assert row is not None and row.downloaded_at is not None
        assert row.expires_at is not None and row.expires_at - row.finished_at == timedelta(hours=24)  # type: ignore[operator]


async def test_expired_export_is_410_and_purged(rorg: ROrg, monkeypatch: pytest.MonkeyPatch) -> None:
    await build_march(rorg)
    export_id = await _async_export(rorg, rorg.admin, monkeypatch)
    async with tenant_session(rorg.tenant_id) as db:
        await db.execute(
            update(Export)
            .where(Export.id == uuid.UUID(export_id))
            .values(expires_at=datetime.now(UTC) - timedelta(minutes=1))
        )
    resp = await rorg.admin.client.get(f"{EXPORTS}/{export_id}/download")
    assert resp.status_code == 410
    assert resp.json()["error"]["code"] == "export_expired"
    from app.tasks import reporting as reporting_tasks

    assert await reporting_tasks._purge_exports(rorg.tenant_id) == 1
    async with tenant_session(rorg.tenant_id) as db:
        row = await db.get(Export, uuid.UUID(export_id))
        assert row is not None and row.status == "expired"
        deleted = (
            await db.execute(text("SELECT deleted_at FROM stored_files WHERE id = :f"), {"f": row.file_id})
        ).scalar_one()
    assert deleted is not None


async def test_permission_revoked_before_async_run(
    app: Any, rorg: ROrg, monkeypatch: pytest.MonkeyPatch
) -> None:
    await build_march(rorg)
    remi = await reporter(app, rorg)
    queued: list[tuple[Any, ...]] = []
    import app.tasks.runtime as runtime

    monkeypatch.setattr(runtime, "enqueue", lambda task, *args, **kw: queued.append(args))
    try:
        export_id = await _async_export(rorg, remi, monkeypatch)
        assert queued == [(str(rorg.tenant_id), export_id)]
        async with tenant_session(rorg.tenant_id) as db:
            role = (await db.execute(select(Role).where(Role.key.like("reporter_%")))).scalar_one()
            role.permissions = {k: v for k, v in role.permissions.items() if k != "report:export"}
        from app.tasks import reporting as reporting_tasks

        assert await reporting_tasks._run_export(rorg.tenant_id, uuid.UUID(export_id)) == "failed"
        # Idempotent: a second run is a no-op.
        assert await reporting_tasks._run_export(rorg.tenant_id, uuid.UUID(export_id)) == "failed"
        async with tenant_session(rorg.tenant_id) as db:
            row = await db.get(Export, uuid.UUID(export_id))
        assert row is not None and (row.status, row.error_code, row.file_id) == ("failed", "forbidden", None)
    finally:
        await remi.client.aclose()


async def test_validation_errors(rorg: ROrg) -> None:
    bad_col = await post_export(rorg.admin, "submissions", "csv", columns=["human_id", "nope"])
    assert bad_col.status_code == 422 and bad_col.json()["error"]["code"] == "unknown_columns"
    hidden_col = await post_export(rorg.admin, "applications", "csv", columns=["candidate_id"])
    assert hidden_col.status_code == 422
    bad_filter = await post_export(rorg.admin, "submissions", "csv", {"nope": 1})
    assert bad_filter.status_code == 422 and bad_filter.json()["error"]["code"] == "unknown_filters"
    bad_value = await post_export(rorg.admin, "submissions", "csv", {"feedback": "meh"})
    assert bad_value.status_code == 422 and bad_value.json()["error"]["code"] == "invalid_filters"
    bad_kind = await post_export(rorg.admin, "notes", "csv")
    assert bad_kind.status_code == 422


async def test_rate_limit_ten_per_hour(rorg: ROrg, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(get_settings(), "rate_limit_enabled", True)
    for _ in range(10):
        assert (await post_export(rorg.admin, "candidates", "csv")).status_code == 200
    resp = await post_export(rorg.admin, "candidates", "csv")
    assert resp.status_code == 429
    assert resp.json()["error"]["code"] == "rate_limited"


async def test_cross_tenant_export_ids_are_404(
    rorg: ROrg, other_rorg: ROrg, monkeypatch: pytest.MonkeyPatch
) -> None:
    await build_march(rorg)
    export_id = await _async_export(rorg, rorg.admin, monkeypatch)
    for url in (f"{EXPORTS}/{export_id}", f"{EXPORTS}/{export_id}/download"):
        assert (await other_rorg.admin.client.get(url)).status_code == 404
    assert (await other_rorg.admin.client.get(EXPORTS)).json()["total"] == 0


async def test_timestamps_use_tenant_timezone() -> None:
    from zoneinfo import ZoneInfo

    from app.services.reporting.lists import COLUMN_INDEX

    col = COLUMN_INDEX["submissions"]["sent_at"]
    values = export_engine.row_values([col], {"sent_at": ist(2026, 3, 9, 0, 30)}, ZoneInfo("Asia/Kolkata"))
    assert values == ["2026-03-09 00:30"]
