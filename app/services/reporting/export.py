r"""CSV / XLSX export engine (Epic 11, A5, FR-61).

Flow of ``POST /exports``: ``report:export`` (router) + the kind's read permission -> filters validated with
the kind's filter model -> columns must be known -> compensation columns **removed** without
``compensation:view`` (listed in ``masked_columns``) -> the kind's ``ListQuery`` (the list endpoint's own
query) is counted:

* ``> EXPORT_MAX_ROWS`` -> 422 ``export_too_large`` (nothing recorded, never truncated);
* ``<= EXPORT_SYNC_MAX_ROWS`` -> streamed now (CSV: server-side cursor, UTF-8 BOM; XLSX: openpyxl
  write-only into a spooled file); an ``exports`` row is recorded ``done`` (no file);
* otherwise -> ``queued`` + ``ExportQueued`` -> ``reporting.run_export`` rebuilds the requester's principal,
  re-checks permissions and masking, writes the file through ``core.storage`` (scanned, ``files`` row,
  24 h expiry) and emits ``ExportCompleted`` (Epic 10's ``export_ready`` notification).

Every export writes ``export.run`` (and every download ``export.downloaded``) in the same transaction as the
``exports`` row (AD-5). 10 exports per user per hour (429); 2 concurrent async exports per tenant.

Formatting (shared by both paths): header = column labels; timestamps in the requester's timezone as
``yyyy-mm-dd HH:MM``; money as two columns (amount, currency); enum labels; lists joined with ``; ``;
formula-injection guard (``= + - @ \\t \\r``): CSV cells get a leading ``'``, XLSX cells are explicit strings.
"""

import csv
import dataclasses
import io
import tempfile
import uuid
from collections.abc import AsyncIterator, Iterable, Sequence
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any
from zoneinfo import ZoneInfo

from loguru import logger
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import audit, rate_limit
from app.core.context import RequestContext
from app.core.database import tenant_session
from app.core.errors import ConflictError, Gone, NotFound, ValidationFailed
from app.core.events import emit
from app.core.pagination import Page, PageParams, paginate
from app.core.permissions import Principal, check
from app.core.settings import get_settings
from app.events.reporting import ExportCompleted, ExportQueued
from app.models.reporting import Export
from app.schemas.reporting import ExportCreate, ExportOut
from app.services.reporting import lists
from app.services.reporting.common import now, principal_of, tenant_of

EXPORT_PERMISSION = "report:export"
BOM = "﻿"
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
BATCH = 1000
WIDTH_SAMPLE_ROWS = 200


# ================================================================ formatting
def headers(columns: Sequence[lists.Column]) -> list[str]:
    out: list[str] = []
    for col in columns:
        out.append(col.label)
        if col.type == "money":
            out.append(f"{col.label} currency")
    return out


def _cell(col: lists.Column, value: Any, tz: ZoneInfo) -> Any:
    if value is None:
        return None
    if col.type == "enum":
        return col.enum_label(value)
    if isinstance(value, datetime):
        return value.astimezone(tz).strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, list | tuple):
        return "; ".join(str(v) for v in value if v is not None)
    if isinstance(value, uuid.UUID):
        return str(value)
    return value


def row_values(columns: Sequence[lists.Column], row: Any, tz: ZoneInfo) -> list[Any]:
    """One export row: typed values (Decimal / int / str / None), money as (amount, currency)."""
    out: list[Any] = []
    for col in columns:
        out.append(_cell(col, row.get(col.key), tz))
        if col.type == "money":
            out.append(row.get(f"{col.key}__currency"))
    return out


def is_formula_like(value: Any) -> bool:
    return isinstance(value, str) and value.startswith(FORMULA_PREFIXES)


def csv_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value, "f")
    text = str(value)
    return "'" + text if is_formula_like(text) else text


def csv_line(values: Iterable[Any]) -> str:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\r\n").writerow([csv_cell(v) for v in values])
    return buf.getvalue()


class XlsxWriter:
    """openpyxl write-only workbook: bold frozen header, auto-filter, widths from the first 200 rows."""

    def __init__(self, sheet_name: str, header: list[str]) -> None:
        from openpyxl import Workbook

        self.wb = Workbook(write_only=True)
        self.ws = self.wb.create_sheet(title=sheet_name[:31])
        self.header = header
        self.buffer: list[list[Any]] = []
        self.started = False
        self.rows = 0

    def _cell(self, value: Any, *, bold: bool = False) -> Any:
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font

        cell = WriteOnlyCell(self.ws, value=float(value) if isinstance(value, Decimal) else value)
        if isinstance(value, str):
            # Explicit string cell: a value starting with ``=`` is never written as a formula.
            cell.data_type = "s"
        if bold:
            cell.font = Font(bold=True)
        return cell

    def _start(self) -> None:
        from openpyxl.utils import get_column_letter

        widths = [len(h) for h in self.header]
        for row in self.buffer:
            for i, value in enumerate(row):
                widths[i] = max(widths[i], len(str(value)) if value is not None else 0)
        for i, width in enumerate(widths, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, width + 2))
        self.ws.freeze_panes = "A2"
        self.ws.append([self._cell(h, bold=True) for h in self.header])
        for row in self.buffer:
            self.ws.append([self._cell(v) for v in row])
        self.buffer = []
        self.started = True

    def append(self, values: list[Any]) -> None:
        self.rows += 1
        if self.started:
            self.ws.append([self._cell(v) for v in values])
            return
        self.buffer.append(values)
        if len(self.buffer) >= WIDTH_SAMPLE_ROWS:
            self._start()

    def close(self) -> bytes:
        from openpyxl.utils import get_column_letter

        if not self.started:
            self._start()
        last = get_column_letter(max(1, len(self.header)))
        self.ws.auto_filter.ref = f"A1:{last}{self.rows + 1}"
        with tempfile.SpooledTemporaryFile(max_size=8 * 1024 * 1024) as fh:
            self.wb.save(fh)
            fh.seek(0)
            return fh.read()


async def render(
    db: AsyncSession, query: lists.ListQuery, fmt: str, tz: ZoneInfo, *, max_rows: int | None = None
) -> tuple[bytes, int]:
    """Stream the query's rows into a CSV / XLSX file in memory; returns (bytes, row count)."""
    columns = query.columns
    result = await db.stream(query.rows.execution_options(yield_per=BATCH))
    count = 0
    if fmt == "csv":
        parts = [BOM, csv_line(headers(columns))]
        async for row in result.mappings():
            count += 1
            parts.append(csv_line(row_values(columns, row, tz)))
            if max_rows is not None and count > max_rows:
                break
        return "".join(parts).encode("utf-8"), count
    writer = XlsxWriter(query.kind, headers(columns))
    async for row in result.mappings():
        count += 1
        writer.append(row_values(columns, row, tz))
        if max_rows is not None and count > max_rows:
            break
    return writer.close(), count


async def csv_stream(tenant_id: uuid.UUID, query: lists.ListQuery, tz: ZoneInfo) -> AsyncIterator[bytes]:
    """Sync CSV body: its own tenant session (the request's has closed) and a server-side cursor."""
    yield (BOM + csv_line(headers(query.columns))).encode("utf-8")
    async with tenant_session(tenant_id) as db:
        result = await db.stream(query.rows.execution_options(yield_per=BATCH))
        chunk: list[str] = []
        async for row in result.mappings():
            chunk.append(csv_line(row_values(query.columns, row, tz)))
            if len(chunk) >= BATCH:
                yield "".join(chunk).encode("utf-8")
                chunk = []
        if chunk:
            yield "".join(chunk).encode("utf-8")


# ================================================================ request
@dataclasses.dataclass(slots=True)
class ExportResult:
    export: Export
    mode: str  # sync | async
    masked_columns: list[str]
    filename: str
    content_type: str
    body: bytes | None = None  # sync XLSX
    stream: AsyncIterator[bytes] | None = None  # sync CSV


def _too_large(row_count: int, maximum: int) -> ValidationFailed:
    return ValidationFailed(
        f"This view has {row_count:,} rows; exports are limited to {maximum:,}. Narrow your filters.",
        code="export_too_large",
        details={"row_count": row_count, "max": maximum},
    )


def filename_for(kind: str, fmt: str, tz: ZoneInfo, at: datetime | None = None) -> str:
    return f"{kind}-{(at or now()).astimezone(tz):%Y-%m-%d}.{fmt}"


def _filters_json(model: Any) -> dict[str, Any]:
    return dict(model.model_dump(mode="json", by_alias=True, exclude_none=True, exclude_defaults=True))


async def _audit_run(db: AsyncSession, row: Export, mode: str) -> None:  # NOSONAR
    audit.record(
        "export.run",
        row.kind,
        row.id,
        session=db,
        export_id=str(row.id),
        format=row.format,
        filters=row.filters,
        columns=row.columns,
        masked_columns=row.masked_columns,
        row_count=row.row_count,
        mode=mode,
    )


async def count_rows(db: AsyncSession, query: lists.ListQuery) -> int:
    """``SELECT count(*)`` of the export's list query (a seam for the size-limit tests)."""
    return int((await db.execute(query.count())).scalar_one())


async def request_export(ctx: RequestContext, db: AsyncSession, body: ExportCreate) -> ExportResult:
    principal = principal_of(ctx)
    check(principal, EXPORT_PERMISSION)
    kind, fmt = body.kind, body.format
    lists.check_kind(principal, kind)
    filters = lists.parse_filters(kind, body.filters)
    query = await lists.build(ctx, db, kind, filters, columns=body.columns, export=True)
    settings = get_settings()
    total = await count_rows(db, query)
    if total > settings.export_max_rows:
        raise _too_large(total, settings.export_max_rows)
    await rate_limit.enforce("export", str(principal.id), settings.export_rate_per_user_hour, 3600)
    mode = "sync" if total <= settings.export_sync_max_rows else "async"
    at = now()
    row = Export(
        id=uuid.uuid4(),
        tenant_id=tenant_of(db),
        requested_by=principal.id,
        kind=kind,
        format=fmt,
        mode=mode,
        filters=_filters_json(filters),
        columns=[c.key for c in query.columns],
        masked_columns=list(query.masked_columns),
        status="done" if mode == "sync" else "queued",
        row_count=total,
        started_at=at if mode == "sync" else None,
        finished_at=at if mode == "sync" else None,
    )
    db.add(row)
    await db.flush()
    await _audit_run(db, row, mode)
    filename = filename_for(kind, fmt, query.tz, at)
    result = ExportResult(
        export=row,
        mode=mode,
        masked_columns=list(query.masked_columns),
        filename=filename,
        content_type=CONTENT_TYPES[fmt],
    )
    if mode == "async":
        emit(db, ExportQueued(tenant_id=row.tenant_id, export_id=row.id, requested_by=principal.id))
        return result
    if fmt == "csv":
        result.stream = csv_stream(row.tenant_id, query, query.tz)
    else:
        result.body, _ = await render(db, query, fmt, query.tz)
    return result


# ================================================================ async run (Celery)
async def _requester_ctx(db: AsyncSession, row: Export) -> RequestContext | None:
    from app.services import entity_refs

    principal = await entity_refs.principal_for_user(db, row.tenant_id, row.requested_by)
    if principal is None:
        return None
    return RequestContext(tenant_id=row.tenant_id, principal=principal)


def _allowed(principal: Principal, kind: str) -> bool:
    return principal.has(EXPORT_PERMISSION) and lists.can_read_kind(principal, kind)


async def run_export(db: AsyncSession, export_id: uuid.UUID) -> str:
    """Produce an async export (idempotent: only a ``queued`` export runs). Returns the final status."""
    row = (
        await db.execute(select(Export).where(Export.id == export_id).with_for_update())
    ).scalar_one_or_none()
    if row is None:
        return "missing"
    if row.status != "queued":
        return row.status
    row.status = "running"
    row.started_at = now()
    await db.flush()
    ctx = await _requester_ctx(db, row)
    if ctx is None or not _allowed(ctx.principal, row.kind):
        return await _fail(db, row, "forbidden")
    try:
        filters = lists.parse_filters(row.kind, row.filters)
        requested = list(row.columns) + [c for c in row.masked_columns if c not in row.columns]
        query = await lists.build(ctx, db, row.kind, filters, columns=requested, export=True)
        settings = get_settings()
        data, count = await render(db, query, row.format, query.tz, max_rows=settings.export_max_rows)
        if count > settings.export_max_rows:
            return await _fail(db, row, "export_too_large")
        from app.services import files_service

        filename = filename_for(row.kind, row.format, query.tz)
        stored = await files_service.store_generated(
            ctx,
            db,
            data,
            key=f"t/{row.tenant_id}/exports/{row.id}.{row.format}",
            filename=filename,
            content_type=CONTENT_TYPES[row.format],
            owner_entity_type="export",
            owner_entity_id=row.id,
        )
    except ValidationFailed as exc:
        logger.warning("Export {} failed: {}", row.id, exc.code)
        return await _fail(db, row, exc.code[:50])
    finished = now()
    row.columns = [c.key for c in query.columns]
    row.masked_columns = list(query.masked_columns)
    row.row_count = count
    row.file_id = stored.id
    row.status = "done"
    row.finished_at = finished
    row.expires_at = finished + timedelta(hours=get_settings().export_file_ttl_hours)
    await db.flush()
    emit(
        db,
        ExportCompleted(
            tenant_id=row.tenant_id,
            export_id=row.id,
            requested_by=row.requested_by,
            kind=row.kind,
            format=row.format,
            row_count=count,
        ),
    )
    return "done"


async def _fail(db: AsyncSession, row: Export, code: str) -> str:
    row.status = "failed"
    row.error_code = code
    row.finished_at = now()
    await db.flush()
    audit.record("export.failed", row.kind, row.id, session=db, export_id=str(row.id), error_code=code)
    return "failed"


async def mark_failed(db: AsyncSession, export_id: uuid.UUID, code: str) -> None:
    row = await db.get(Export, export_id, with_for_update=True)
    if row is not None and row.status in ("queued", "running"):
        await _fail(db, row, code)


# ================================================================ my exports / download / cleanup
def export_out(row: Export) -> ExportOut:
    downloadable = row.status == "done" and row.file_id is not None
    return ExportOut(
        id=row.id,
        kind=row.kind,
        format=row.format,
        mode=row.mode,
        status=row.status,
        filters=row.filters,
        columns=list(row.columns),
        masked_columns=list(row.masked_columns),
        row_count=row.row_count,
        error_code=row.error_code,
        created_at=row.created_at,
        started_at=row.started_at,
        finished_at=row.finished_at,
        expires_at=row.expires_at,
        downloaded_at=row.downloaded_at,
        download_url=f"/api/v1/exports/{row.id}/download" if downloadable else None,
    )


async def _mine(
    ctx: RequestContext, db: AsyncSession, export_id: uuid.UUID, *, for_update: bool = False
) -> Export:
    principal = principal_of(ctx)
    stmt = select(Export).where(Export.id == export_id, Export.requested_by == principal.id)
    if for_update:
        stmt = stmt.with_for_update()
    row = (await db.execute(stmt)).scalar_one_or_none()
    if row is None:
        raise NotFound("Export not found")
    return row


async def list_exports(ctx: RequestContext, db: AsyncSession, params: PageParams) -> Page[ExportOut]:
    principal = principal_of(ctx)
    stmt = (
        select(Export)
        .where(Export.requested_by == principal.id)
        .order_by(Export.created_at.desc(), Export.id.desc())
    )
    page = await paginate(db, stmt, params)
    return Page[ExportOut](
        items=[export_out(r) for r in page.items], total=page.total, page=page.page, page_size=page.page_size
    )


async def get_export(ctx: RequestContext, db: AsyncSession, export_id: uuid.UUID) -> ExportOut:
    return export_out(await _mine(ctx, db, export_id))


def _expired(row: Export, at: datetime) -> bool:
    return row.status == "expired" or (row.expires_at is not None and row.expires_at <= at)


async def download(ctx: RequestContext, db: AsyncSession, export_id: uuid.UUID) -> uuid.UUID:
    """Requester-only download of an async export: returns the stored file id (``/files/{id}/download``
    re-checks access). 404 for anyone else, 410 ``export_expired`` after expiry, 409 while not ready."""
    principal = principal_of(ctx)
    check(principal, EXPORT_PERMISSION)
    row = await _mine(ctx, db, export_id, for_update=True)
    if _expired(row, now()):
        raise Gone("This export has expired", code="export_expired")
    if row.status != "done" or row.file_id is None:
        raise ConflictError("This export is not ready for download", code="export_not_ready")
    row.downloaded_at = now()
    await db.flush()
    audit.record(
        "export.downloaded", row.kind, row.id, session=db, export_id=str(row.id), file_id=str(row.file_id)
    )
    return row.file_id


async def file_access(ctx: RequestContext, db: AsyncSession, export_id: uuid.UUID, write: bool) -> bool:
    """``files`` access hook for export files: the requester only, with ``report:export``, until expiry."""
    principal = ctx.principal
    if write or not isinstance(principal, Principal) or principal.type != "user":
        return False
    if not principal.has(EXPORT_PERMISSION):
        return False
    row = await db.get(Export, export_id)
    return (
        row is not None
        and row.requested_by == principal.id
        and row.status == "done"
        and not _expired(row, now())
    )


async def purge_expired(db: AsyncSession, at: datetime | None = None) -> int:
    """Delete the files of expired exports and mark them ``expired`` (one tenant session)."""
    from app.services import files_service

    at = at or now()
    rows = (
        (
            await db.execute(
                select(Export)
                .where(Export.status == "done", Export.expires_at.is_not(None), Export.expires_at <= at)
                .with_for_update(skip_locked=True)
            )
        )
        .scalars()
        .all()
    )
    for row in rows:
        if row.file_id is not None:
            await files_service.purge_file(db, row.file_id)
        row.status = "expired"
    await db.flush()
    return len(rows)
