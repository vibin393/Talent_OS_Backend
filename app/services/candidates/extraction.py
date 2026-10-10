"""Resume text extraction (FR-19): ``pypdf`` for PDF, ``python-docx`` for DOCX. No OCR in Phase 1.

* A PDF with fewer than 50 non-whitespace characters is ``image_only`` ("parse failed - manual entry").
* An encrypted PDF is ``encrypted``; anything unreadable is ``failed``.
"""

import dataclasses
import io
import re
from typing import Literal

from loguru import logger

TextStatus = Literal["ok", "image_only", "encrypted", "failed"]

MIN_TEXT_CHARS = 50
MAX_TEXT_CHARS = 200_000
_NON_WS = re.compile(r"\S")


@dataclasses.dataclass(frozen=True, slots=True)
class Extraction:
    status: TextStatus
    text: str | None


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()[:MAX_TEXT_CHARS]


def _pdf(data: bytes) -> Extraction:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                # Owner-password-only PDFs open with an empty user password.
                if not reader.decrypt(""):
                    return Extraction("encrypted", None)
            except Exception:
                return Extraction("encrypted", None)
        parts = [(page.extract_text() or "") for page in reader.pages]
    except PdfReadError:
        return Extraction("failed", None)
    except Exception as exc:  # malformed files raise all sorts of errors
        logger.info("PDF text extraction failed: {}", type(exc).__name__)
        return Extraction("failed", None)
    text = _clean("\n".join(parts))
    if len(_NON_WS.findall(text)) < MIN_TEXT_CHARS:
        return Extraction("image_only", text or None)
    return Extraction("ok", text)


def _docx(data: bytes) -> Extraction:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
    except Exception as exc:
        logger.info("DOCX text extraction failed: {}", type(exc).__name__)
        return Extraction("failed", None)
    text = _clean("\n".join(parts))
    if not _NON_WS.search(text):
        return Extraction("failed", None)
    return Extraction("ok", text)


def extract_text(data: bytes, kind: str) -> Extraction:
    if kind == "pdf":
        return _pdf(data)
    if kind == "docx":
        return _docx(data)
    return Extraction("failed", None)
