# ruff: file-ignore[line-too-long]
r"""Item import (Epic 40 A3): CSV / XLSX -> preview with per-row errors -> commit; and the go-live import function
``import_assessment_item`` (FR-263, foundation 0.7).

Columns (one row per item): ``kind, stem, points, difficulty, tags, skill_keys, language, code, options, answer,
rubric``. Lists are ``;``-separated. Per kind:

* ``single`` / ``true_false`` / ``multi``: ``options`` = choice texts; ``answer`` = 1-based choice number(s) (``2`` or ``1;3``);
* ``numeric``: ``answer`` = ``value`` or ``value|tolerance_abs``;
* ``short_text``: ``answer`` = accepted answers;
* ``code_output``: ``answer`` = the expected output (``\\n`` for new lines);
* ``code_fill``: ``answer`` = ``1=answer;2=answer`` for the ``{{blank:n}}`` markers in ``code``;
* ``code_order``: ``options`` = the code lines in the correct order (shuffled per attempt at delivery);
* ``code_find_bug``: ``answer`` = buggy line numbers;
* ``code_freeform`` / ``essay``: ``rubric`` = ``label:max_points`` criteria.

Deviation (recorded): batches live in the domain table ``assessment_item_imports`` -- the P1 ``import_batches`` /
``import_items`` tables are resume-import specific (file / candidate columns, no ``kind``).
"""

import csv
import io
import uuid
from decimal import Decimal, InvalidOperation
from typing import Any

from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.context import RequestContext
from app.core.errors import DomainError, NotFound
from app.models.assessments import ITEM_KINDS, AssessmentItemImport
from app.schemas.assessments import (
    AssessmentItemImportIn,
    ImportCommitOut,
    ImportPreviewOut,
    ImportRowOut,
    ItemCreate,
)
from app.services.assessments import items
from app.services.assessments.common import conflict, invalid, now, tenant_of, user_id

COLUMNS = (
    "kind",
    "stem",
    "points",
    "difficulty",
    "tags",
    "skill_keys",
    "language",
    "code",
    "options",
    "answer",
    "rubric",
)
MAX_ROWS = 1000


def _split(value: str | None) -> list[str]:
    return [p.strip() for p in (value or "").split(";") if p.strip()]


def _letters(n: int) -> list[str]:
    return [chr(ord("a") + i) for i in range(n)]


def _indices(value: str | None, n: int) -> list[int]:
    out: list[int] = []
    for p in _split(value):
        if not p.isdigit() or not 1 <= int(p) <= n:
            msg = f"answer '{p}' is not a choice number between 1 and {n}"
            raise ValueError(msg)
        out.append(int(p) - 1)
    return out


def to_item(row: AssessmentItemImportIn) -> ItemCreate:
    """Map one import row to an ``ItemCreate`` (ValueError with a readable message on bad input)."""
    kind = row.kind
    options: dict[str, Any] = {}
    key: dict[str, Any] | None = None
    rubric: dict[str, Any] | None = None
    if kind in ("single", "multi", "true_false"):
        texts = _split(row.options)
        ids = _letters(len(texts))
        options = {"choices": [{"id": i, "text": t} for i, t in zip(ids, texts, strict=True)]}
        picked = _indices(row.answer, len(texts))
        if not picked:
            raise ValueError("answer is required")
        key = {"correct": [ids[i] for i in picked]} if kind == "multi" else {"correct": ids[picked[0]]}
    elif kind == "numeric":
        raw = (row.answer or "").split("|")
        try:
            key = {"value": str(Decimal(raw[0].strip()))}
            if len(raw) > 1 and raw[1].strip():
                key["tolerance_abs"] = str(Decimal(raw[1].strip()))
        except InvalidOperation as exc:
            raise ValueError("answer must be a number (optionally value|tolerance)") from exc
    elif kind == "short_text":
        key = {"accepted": _split(row.answer)}
    elif kind == "code_output":
        key = {"accepted": [(row.answer or "").replace("\\n", "\n")]}
    elif kind == "code_fill":
        blanks: dict[str, Any] = {}
        for part in _split(row.answer):
            n, _, ans = part.partition("=")
            if not n.strip().isdigit() or not ans.strip():
                msg = f"blank answer '{part}' must look like 1=answer"
                raise ValueError(msg)
            blanks.setdefault(n.strip(), {"accepted": []})["accepted"].append(ans.strip())
        key = {"blanks": blanks}
    elif kind == "code_order":
        lines = _split(row.options)
        ids = [f"l{i + 1}" for i in range(len(lines))]
        options = {"lines": [{"id": i, "text": t} for i, t in zip(ids, lines, strict=True)]}
        key = {"order": ids, "partial": "lcs"}
    elif kind == "code_find_bug":
        try:
            key = {"bug_lines": [int(p) for p in _split(row.answer)]}
        except ValueError as exc:
            raise ValueError("answer must list line numbers") from exc
    else:
        criteria = []
        for i, part in enumerate(_split(row.rubric)):
            label, _, pts = part.rpartition(":")
            try:
                criteria.append(
                    {
                        "key": f"c{i + 1}",
                        "label": label.strip() or part,
                        "max_points": str(Decimal(pts.strip())),
                    }
                )
            except InvalidOperation as exc:
                msg = f"rubric criterion '{part}' must look like label:points"
                raise ValueError(msg) from exc
        rubric = {"criteria": criteria}
    return ItemCreate(
        kind=kind,
        stem=row.stem,
        points=row.points,
        difficulty=row.difficulty,
        tags=_split(row.tags),
        skill_keys=_split(row.skill_keys),
        language=row.language or None,
        code=(row.code or "").replace("\\n", "\n") or None,
        options=options,
        answer_key=key,
        rubric=rubric,
    )


def _check(data: ItemCreate) -> None:
    items.validate_item(data.kind, data)
    items._skill_keys(data.skill_keys)


def _read_rows(content: bytes, filename: str | None) -> list[dict[str, str]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = wb.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h or "").strip().lower() for h in rows[0]]
        return [
            {
                header[i]: ("" if v is None else str(v))
                for i, v in enumerate(r)
                if i < len(header) and header[i]
            }
            for r in rows[1:]
            if any(v not in (None, "") for v in r)
        ]
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [{(k or "").strip().lower(): (v or "") for k, v in r.items()} for r in reader]


async def preview(
    ctx: RequestContext, db: AsyncSession, content: bytes, filename: str | None
) -> ImportPreviewOut:
    try:
        raw_rows = _read_rows(content, filename)
    except Exception as exc:
        raise invalid("import_unreadable", "The file could not be read as CSV or XLSX", "file") from exc
    if not raw_rows:
        raise invalid("import_empty", "The file has no item rows", "file")
    if len(raw_rows) > MAX_ROWS:
        raise invalid("import_too_large", f"At most {MAX_ROWS} rows per import", "file")
    out_rows: list[ImportRowOut] = []
    stored: list[dict[str, Any]] = []
    for n, raw in enumerate(raw_rows, start=2):
        errors: list[str] = []
        item: ItemCreate | None = None
        kind = (raw.get("kind") or "").strip() or None
        try:
            if kind not in ITEM_KINDS:
                raise ValueError(f"unknown kind '{kind}'")
            values = {k: raw[k] for k in COLUMNS if (raw.get(k) or "").strip()}
            parsed = AssessmentItemImportIn.model_validate(
                values | {"kind": kind, "stem": raw.get("stem") or ""}
            )
            item = to_item(parsed)
            _check(item)
        except ValidationError as exc:
            errors = [f"{'.'.join(str(p) for p in e.get('loc', ()))}: {e.get('msg')}" for e in exc.errors()][
                :5
            ]
        except DomainError as exc:
            errors = [exc.message]
        except ValueError as exc:
            errors = [str(exc)]
        out_rows.append(
            ImportRowOut(
                row=n, kind=kind, stem=(raw.get("stem") or "")[:200] or None, errors=errors, valid=not errors
            )
        )
        stored.append(
            {
                "row": n,
                "errors": errors,
                "item": item.model_dump(mode="json") if item and not errors else None,
            }
        )
    batch = AssessmentItemImport(
        id=uuid.uuid4(),
        tenant_id=tenant_of(db),
        status="previewed",
        filename=(filename or "")[:255] or None,
        rows=stored,
        created_by=user_id(ctx),
    )
    db.add(batch)
    await db.flush()
    valid = sum(1 for r in out_rows if r.valid)
    return ImportPreviewOut(
        batch_id=batch.id, rows=out_rows, valid_count=valid, error_count=len(out_rows) - valid
    )


async def commit(ctx: RequestContext, db: AsyncSession, batch_id: uuid.UUID) -> ImportCommitOut:
    batch = (
        await db.execute(
            select(AssessmentItemImport).where(AssessmentItemImport.id == batch_id).with_for_update()
        )
    ).scalar_one_or_none()
    if batch is None:
        raise NotFound("Import not found")
    if batch.status != "previewed":
        raise conflict("import_committed", "This import was already committed")
    ids: list[uuid.UUID] = []
    for row in batch.rows or []:
        if row.get("item"):
            created = await items.create_item(ctx, db, ItemCreate.model_validate(row["item"]))
            ids.append(created.id)
    batch.status = "committed"
    batch.committed_at = now()
    await db.flush()
    return ImportCommitOut(created=len(ids), item_ids=ids)


async def import_assessment_item(
    ctx: RequestContext, db: AsyncSession, data: AssessmentItemImportIn, *, validate_only: bool = False
) -> uuid.UUID | None:
    """Go-live import of one item row (FR-263). Raises ``ValidationFailed`` naming the field."""
    try:
        item = to_item(data)
    except ValueError as exc:
        raise invalid("invalid_item", str(exc), "answer") from exc
    _check(item)
    if validate_only:
        return None
    return (await items.create_item(ctx, db, item)).id


__all__ = ["COLUMNS", "commit", "import_assessment_item", "preview", "to_item"]
