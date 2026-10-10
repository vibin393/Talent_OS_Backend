"""Epic 28 Payroll Export: periods, previews (no postings), finalize (exactly-once postings, locks,
idempotency, race), supplementary, reopen, files (Decimal strings, empty statutory columns, masking),
mark processed,
human-only, immutability, exceptions and tenant isolation.

Payable items come from a test posting source (``e28_hours`` / ``e28_exp``) so the payroll pipeline is
exercised without driving the timesheet / expense approval flows; the real loaders (Epics 25 / 27) are
covered by their own modules.
"""

import asyncio
import csv
import io
import re
import uuid
from collections.abc import Iterator
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from typing import Any

import pytest
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError

from app.core.context import RequestContext
from app.core.database import tenant_session
from app.core.errors import Forbidden
from app.core.permissions import CATALOG, Principal, Scope
from app.models.legal_entities import LegalEntity
from app.services import postings
from app.services.payroll import earnings, exports, files
from tests.modules.demands.helpers import DOrg, admin_ctx, h
from tests.modules.placements.helpers import WORKERS, direct_placement

P = "/api/v1/payroll"
HOURS, EXP = "e28_hours", "e28_exp"
FAKE: list[postings.SourceItem] = []
BANK = {
    "holder_name": "Asha Rao",
    "account_number": "123456789012",
    "account_number_confirm": "123456789012",
    "ifsc": "HDFC0001234",
}


def _loader_for(kind: str) -> Any:
    async def load(ctx: RequestContext, db: Any, q: postings.SourceQuery) -> list[postings.SourceItem]:
        return [it for it in _matching(q) if it.source_kind == kind]

    return load


def _matching(q: postings.SourceQuery) -> list[postings.SourceItem]:
    out = []
    for it in FAKE:
        if it.item_date > q.period_end or (q.period_start and it.item_date < q.period_start):
            continue
        if q.legal_entity_id and it.attributes.get("legal_entity_id") != q.legal_entity_id:
            continue
        if q.placement_ids is not None and it.placement_id not in q.placement_ids:
            continue
        if q.payee_kind and it.attributes.get("payee") != q.payee_kind:
            continue
        out.append(it)
    return out


@pytest.fixture(autouse=True)
def fake_source(monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    for kind in (HOURS, EXP):
        if kind not in postings.posting_sources():
            postings.register_posting_source(kind, _loader_for(kind), ledgers=frozenset({"payable"}))
    monkeypatch.setattr(earnings, "EXPENSE_KINDS", frozenset({"expense_line", EXP}))
    FAKE.clear()
    yield
    FAKE.clear()


def month_start() -> date:
    return date.today().replace(day=1)


class World:
    def __init__(self, org: DOrg, le_id: uuid.UUID, placement: dict[str, Any]) -> None:
        self.org, self.le_id, self.placement = org, le_id, placement

    @property
    def pid(self) -> uuid.UUID:
        return uuid.UUID(self.placement["id"])

    @property
    def cid(self) -> uuid.UUID:
        return uuid.UUID(self.placement["candidate_id"])

    def hours(
        self, day: int, q: str = "7.83", category: str = "regular", approved_at: datetime | None = None
    ) -> uuid.UUID:
        sid = uuid.uuid4()
        FAKE.append(
            postings.SourceItem(
                HOURS,
                sid,
                self.pid,
                month_start() + timedelta(days=day),
                Decimal(q),
                None,
                "INR",
                {
                    "category": category,
                    "candidate_id": self.cid,
                    "legal_entity_id": self.le_id,
                    "payee": "worker",
                    "approved_at": approved_at or datetime.now(UTC),
                    "unit": "hour",
                },
            )
        )
        return sid

    def expense(self, amount: str = "1026.57") -> uuid.UUID:
        sid = uuid.uuid4()
        FAKE.append(
            postings.SourceItem(
                EXP,
                sid,
                self.pid,
                month_start() + timedelta(days=2),
                Decimal(1),
                Decimal(amount),
                "INR",
                {
                    "category": "travel",
                    "candidate_id": self.cid,
                    "legal_entity_id": self.le_id,
                    "payee": "worker",
                    "reimbursable": True,
                    "approved_at": datetime.now(UTC),
                    "unit": "item",
                },
            )
        )
        return sid

    async def period_id(self) -> str:
        resp = await self.org.admin.client.get(f"{P}/periods", params={"legal_entity_id": str(self.le_id)})
        assert resp.status_code == 200, resp.text
        rows = [r for r in resp.json() if r["period_start"] == month_start().isoformat()]
        assert rows, resp.json()
        return str(rows[0]["id"])


async def le_of(org: DOrg) -> uuid.UUID:
    async with tenant_session(org.tenant_id) as db:
        return (
            await db.execute(select(LegalEntity.id).order_by(LegalEntity.created_at).limit(1))
        ).scalar_one()


async def world(org: DOrg, *, bank: bool = True, **body: Any) -> World:
    le_id = await le_of(org)
    m = org.admin
    resp = await m.client.put(
        f"{P}/calendars/{le_id}",
        json={"cutoff_day": "last", "generate_from": (month_start() - timedelta(days=40)).isoformat()},
        headers=h(m),
    )
    assert resp.status_code == 200, resp.text
    start = (month_start() - timedelta(days=35)).isoformat()
    p = await direct_placement(org, planned_start_date=start, pay_rate="437.5000", **body)
    w = World(org, le_id, p)
    if bank:
        await add_bank(org, w.cid)
    return w


async def add_bank(org: DOrg, cid: uuid.UUID) -> None:
    resp = await org.admin.client.post(f"{WORKERS}/{cid}/bank-accounts", json=BANK, headers=h(org.admin))
    assert resp.status_code in (200, 201), resp.text
    async with tenant_session(org.tenant_id) as db:
        await db.execute(
            text("UPDATE worker_bank_accounts SET status = 'active' WHERE candidate_id = :c"), {"c": cid}
        )


async def preview(w: World, period_id: str, **body: Any) -> dict[str, Any]:
    resp = await w.org.admin.client.post(
        f"{P}/periods/{period_id}/previews", json=body, headers=h(w.org.admin)
    )
    assert resp.status_code == 201, resp.text
    return dict(resp.json())


async def finalize(w: World, version_id: str, key: str | None = None) -> Any:
    headers = {**h(w.org.admin), "Idempotency-Key": key or uuid.uuid4().hex}
    return await w.org.admin.client.post(f"{P}/versions/{version_id}/finalize", headers=headers)


async def payable_count(org: DOrg, **where: Any) -> int:
    sql = "SELECT count(*) FROM payable_postings WHERE state <> 'released'"
    for k in where:
        sql += f" AND {k} = :{k}"
    async with tenant_session(org.tenant_id) as db:
        return int((await db.execute(text(sql), where)).scalar_one())


# ================================================================ periods
async def test_calendar_generates_non_overlapping_periods(dorg: DOrg) -> None:
    le_id = await le_of(dorg)
    resp = await dorg.admin.client.put(
        f"{P}/calendars/{le_id}",
        json={"cutoff_day": 25, "generate_from": "2026-10-10"},
        headers=h(dorg.admin),
    )
    assert resp.status_code == 200 and resp.json()["cutoff_day"] == 25
    periods = (await dorg.admin.client.get(f"{P}/periods", params={"legal_entity_id": str(le_id)})).json()
    spans = {(p["period_start"], p["period_end"]) for p in periods}
    assert ("2026-09-26", "2026-10-25") in spans
    again = await dorg.admin.client.put(
        f"{P}/calendars/{le_id}", json={"cutoff_day": 25}, headers=h(dorg.admin)
    )
    assert again.status_code == 200
    assert len(
        (await dorg.admin.client.get(f"{P}/periods", params={"legal_entity_id": str(le_id)})).json()
    ) == len(periods)
    async with tenant_session(dorg.tenant_id) as db:
        with pytest.raises(DBAPIError, match="ex_pay_periods_no_overlap"):
            await db.execute(
                text(
                    "INSERT INTO pay_periods (id, tenant_id, legal_entity_id, human_id, "
                    "period_start, period_end) "
                    "VALUES (:i, :t, :l, 'X', '2026-10-01', '2026-10-05')"
                ),
                {"i": uuid.uuid4(), "t": dorg.tenant_id, "l": le_id},
            )


# ================================================================ end to end
async def test_preview_finalize_file_processed(dorg: DOrg) -> None:
    w = await world(dorg)
    for d in range(3):
        w.hours(d)
    w.expense()
    period_id = await w.period_id()
    before = await payable_count(dorg)
    p1 = await preview(w, period_id)
    p2 = await preview(w, period_id)
    assert await payable_count(dorg) == before  # previews never post
    first = (await dorg.admin.client.get(f"{P}/versions/{p1['id']}")).json()
    assert first["status"] == "superseded" and p2["status"] == "generated"
    assert p2["totals"]["INR"]["gross_earnings_inputs"] == "10276.89"  # 3 x 3425.63
    rows = (await dorg.admin.client.get(f"{P}/versions/{p2['id']}/rows")).json()
    assert (
        len(rows) == 1
        and rows[0]["earnings_regular"] == "10276.89"
        and rows[0]["reimbursements_total"] == "1026.57"
    )
    assert rows[0]["hours_regular"] == "23.49"

    stale = await finalize(w, p1["id"])
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "preview_stale"

    key = uuid.uuid4().hex
    resp = await finalize(w, p2["id"], key)
    assert resp.status_code == 200, resp.text
    final = resp.json()
    assert final["kind"] == "final" and final["status"] == "finalized" and final["content_hash"]
    assert await payable_count(dorg, document_id=uuid.UUID(final["id"])) == 4  # one posting per item
    replay = await finalize(w, p2["id"], key)
    assert replay.status_code == 200 and replay.json()["id"] == final["id"]
    assert await payable_count(dorg) == before + 4
    period = (await dorg.admin.client.get(f"{P}/periods/{period_id}")).json()
    assert period["status"] == "locked" and period["final_count"] == 1

    # postings equal items (quantity / amount)
    async with tenant_session(dorg.tenant_id) as db:
        pairs = (
            await db.execute(
                text(
                    "SELECT i.amount, p.amount, i.quantity, p.quantity FROM payroll_export_items i "
                    "JOIN payable_postings p ON p.source_id = i.source_id AND p.target_id = i.row_id "
                    "WHERE i.version_id = :v"
                ),
                {"v": uuid.UUID(final["id"])},
            )
        ).all()
    assert len(pairs) == 4 and all(a == b and q1 == q2 for a, b, q1, q2 in pairs)

    # CSV through the default mapping (admin holds identity:view -> unmasked + audited)
    f = await dorg.admin.client.get(f"{P}/versions/{final['id']}/file", params={"format": "csv"})
    assert f.status_code == 200, f.text
    body = f.text
    assert "10276.89" in body and "1026.57" in body and "123456789012" in body
    assert not re.search(r"\de-\d|\d\.\d{3,}", body.replace("23.49", ""))
    reader = list(csv.DictReader(io.StringIO(body)))
    for col in (
        "PF (employee)",
        "PF (employer)",
        "ESI (employee)",
        "ESI (employer)",
        "Professional tax",
        "TDS (salary)",
        "Other deductions",
    ):
        assert all(r[col] == "" for r in reader), col
    async with tenant_session(dorg.tenant_id) as db:
        n = (
            await db.execute(text("SELECT count(*) FROM audit_logs WHERE action = 'payroll.file_unmasked'"))
        ).scalar_one()
    assert n >= 1

    # masked without identity:view
    perms = {k: Scope.TENANT for k in CATALOG if k != "identity:view"}
    ctx = RequestContext(
        tenant_id=dorg.tenant_id,
        principal=Principal(type="user", id=dorg.admin.user_id, tenant_id=dorg.tenant_id, permissions=perms),
    )
    async with tenant_session(dorg.tenant_id) as db:
        masked = await files.build_file(ctx, db, uuid.UUID(final["id"]), None, "csv")
        xlsx = await files.build_file(ctx, db, uuid.UUID(final["id"]), None, "xlsx")
    assert b"XXXXXXXX9012" in masked.data and b"123456789012" not in masked.data
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(xlsx.data)).active
    header = [c.value for c in ws[1]]
    row2 = {header[i]: c for i, c in enumerate(ws[2])}
    assert (
        row2["Regular earnings"].number_format == "0.00" and str(row2["Regular earnings"].value) == "10276.89"
    )
    assert row2["PF (employee)"].value in (None, "")

    # mark processed twice -> one transition
    for _ in range(2):
        mp = await dorg.admin.client.post(
            f"{P}/versions/{final['id']}/mark-processed",
            json={"payroll_reference": "RUN-10"},
            headers=h(dorg.admin),
        )
        assert mp.status_code == 200 and mp.json()["status"] == "processed"
    async with tenant_session(dorg.tenant_id) as db:
        n = (
            await db.execute(text("SELECT count(*) FROM audit_logs WHERE action = 'payroll.processed'"))
        ).scalar_one()
    assert n == 1

    # immutable at the database once final
    async with tenant_session(dorg.tenant_id) as db:
        with pytest.raises(DBAPIError, match="document_immutable"):
            await db.execute(
                text("UPDATE payroll_export_rows SET gross_earnings_inputs = 1 WHERE version_id = :v"),
                {"v": uuid.UUID(final["id"])},
            )
    async with tenant_session(dorg.tenant_id) as db:
        with pytest.raises(DBAPIError, match="document_immutable"):
            await db.execute(
                text("DELETE FROM payroll_export_versions WHERE id = :v"), {"v": uuid.UUID(final["id"])}
            )
    gone = await dorg.admin.client.delete(f"{P}/versions/{final['id']}", headers=h(dorg.admin))
    assert gone.status_code == 409


async def test_supplementary_reopen_and_nothing_to_export(dorg: DOrg) -> None:
    w = await world(dorg)
    w.hours(0)
    period_id = await w.period_id()
    pv = await preview(w, period_id)
    final = (await finalize(w, pv["id"])).json()
    empty = await w.org.admin.client.post(f"{P}/periods/{period_id}/previews", json={}, headers=h(dorg.admin))
    assert empty.status_code == 409 and empty.json()["error"]["code"] == "nothing_to_export"
    # late approval (adjustment) -> supplementary only
    late = w.hours(1, "-1.00", "ot")
    sp = await preview(w, period_id)
    assert sp["supplementary"] is True and sp["item_count"] == 1
    resp = await finalize(w, sp["id"])
    assert resp.status_code == 200, resp.text
    assert resp.json()["kind"] == "supplementary"
    assert await payable_count(dorg, source_id=late) == 1
    # reopen (human, reason) leaves the final byte-identical
    ro = await dorg.admin.client.post(
        f"{P}/periods/{period_id}/reopen", json={"reason": "Missed bonus hours"}, headers=h(dorg.admin)
    )
    assert ro.status_code == 200 and ro.json()["status"] == "open"
    again = (await dorg.admin.client.get(f"{P}/versions/{final['id']}")).json()
    assert again["content_hash"] == final["content_hash"] and again["status"] == "finalized"
    w.hours(2)
    sp2 = await preview(w, period_id)
    assert sp2["supplementary"] is True
    assert (await finalize(w, sp2["id"])).json()["kind"] == "supplementary"


async def test_stale_preview_and_vendor_posted_item_excluded(dorg: DOrg) -> None:
    w = await world(dorg)
    w.hours(0)
    taken = w.hours(1)
    async with tenant_session(dorg.tenant_id) as db:
        await postings.record_posted(
            admin_ctx(dorg),
            db,
            "payable",
            [
                postings.PostingLine(
                    HOURS,
                    taken,
                    "vendor_statement_line",
                    uuid.uuid4(),
                    None,
                    Decimal("7.83"),
                    Decimal("1.00"),
                    "INR",
                    w.pid,
                    "vendor",
                )
            ],
        )
    period_id = await w.period_id()
    pv = await preview(w, period_id)
    assert pv["item_count"] == 1  # the vendor-stated item is never collected
    w.hours(2)  # new approval after the preview
    resp = await finalize(w, pv["id"])
    assert resp.status_code == 409 and resp.json()["error"]["code"] == "preview_stale"


async def test_concurrent_finalize_exactly_one_wins(dorg: DOrg) -> None:
    w = await world(dorg)
    w.hours(0)
    w.hours(1)
    pv = await preview(w, await w.period_id())
    before = await payable_count(dorg)
    a, b = await asyncio.gather(finalize(w, pv["id"]), finalize(w, pv["id"]))
    codes = sorted([a.status_code, b.status_code])
    assert codes == [200, 409], (a.text, b.text)
    assert await payable_count(dorg) == before + 2


async def test_human_only_finalize_and_reopen(dorg: DOrg) -> None:
    w = await world(dorg)
    w.hours(0)
    period_id = await w.period_id()
    pv = await preview(w, period_id)
    for kind in ("workflow", "ai", "system", "api_key"):
        ctx = RequestContext(tenant_id=dorg.tenant_id, actor_kind=kind)
        async with tenant_session(dorg.tenant_id) as db:
            with pytest.raises(Forbidden):
                await exports.finalize(ctx, db, uuid.UUID(pv["id"]))
            with pytest.raises(Forbidden):
                from app.services.payroll import pay_periods

                await pay_periods.reopen(ctx, db, uuid.UUID(period_id), "automation")


async def test_missing_bank_details_exception_and_exclusion(dorg: DOrg) -> None:
    w = await world(dorg, bank=False)
    w.hours(0)
    period_id = await w.period_id()
    exc = (await dorg.admin.client.get(f"{P}/periods/{period_id}/exceptions")).json()
    assert exc["blocking"] is True and exc["counts"]["missing_bank_details"] == 1
    blocked = await dorg.admin.client.post(
        f"{P}/periods/{period_id}/previews", json={}, headers=h(dorg.admin)
    )
    assert blocked.status_code == 409 and blocked.json()["error"]["code"] == "payroll_exceptions"
    pv = await preview(w, period_id, exclude_exceptions=True)
    assert pv["item_count"] == 0 and pv["exceptions"]["missing_bank_details"]
    assert (await finalize(w, pv["id"])).json()["error"]["code"] == "nothing_to_export"


async def test_mappings_validation(dorg: DOrg) -> None:
    m = dorg.admin
    listed = (await m.client.get(f"{P}/mappings")).json()
    assert any(x["is_default"] for x in listed)
    bad = await m.client.post(
        f"{P}/mappings",
        json={"name": "X", "columns": [{"source_field": "net_pay", "header": "Net"}]},
        headers=h(m),
    )
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "unknown_source_field"
    no_id = await m.client.post(
        f"{P}/mappings",
        json={"name": "X", "columns": [{"source_field": "hours_regular", "header": "Hrs"}]},
        headers=h(m),
    )
    assert no_id.status_code == 422 and no_id.json()["error"]["code"] == "identity_column_required"
    ok = await m.client.post(
        f"{P}/mappings",
        json={
            "name": "Keka",
            "format": "xlsx",
            "columns": [
                {"source_field": "worker_human_id", "header": "Emp"},
                {"source_field": "pf_employee", "header": "PF"},
            ],
        },
        headers=h(m),
    )
    assert ok.status_code == 201, ok.text


async def test_tenant_isolation(dorg: DOrg, other_dorg: DOrg) -> None:
    w = await world(dorg)
    w.hours(0)
    period_id = await w.period_id()
    pv = await preview(w, period_id)
    o = other_dorg.admin
    for path in (
        f"/periods/{period_id}",
        f"/periods/{period_id}/exceptions",
        f"/versions/{pv['id']}",
        f"/versions/{pv['id']}/rows",
        f"/versions/{pv['id']}/file",
    ):
        assert (await o.client.get(f"{P}{path}")).status_code == 404, path
    resp = await o.client.post(
        f"{P}/versions/{pv['id']}/finalize", headers={**h(o), "Idempotency-Key": uuid.uuid4().hex}
    )
    assert resp.status_code == 404
    assert (
        await o.client.post(f"{P}/periods/{period_id}/previews", json={}, headers=h(o))
    ).status_code == 404
