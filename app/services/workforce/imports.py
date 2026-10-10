# ruff: noqa: D401, E501
"""Plan line import (Epic 44, FR-222) and the go-live importer (foundation 0.7, FR-263).

``import_lines``: a CSV or XLSX file (header-mapped columns ``role, skills, headcount, type, location, start_month,
budget, bill_rate, unit, notes``; skills ``"Java:4; Spring:3"``) is validated row by row
(``validation.parse_import_row`` + ``validate_line``) and recorded through ``imports_service.register_preview`` with
the per-row errors. ``apply=true`` inserts every valid row in one transaction (the batch is marked applied; the
same file cannot be applied twice to the same plan). ``import_plan_line`` is the go-live entry point: one row,
the plan named by ``plan_name`` (an editable plan).
"""

import csv
import hashlib
import io
import uuid
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import audit
from app.core.context import RequestContext
from app.core.errors import ValidationFailed
from app.models.workforce import WorkforcePlan
from app.schemas.workforce import ImportRowError, LineImportOut, PlanLineImportIn
from app.services.workforce.common import EDITABLE, WRITE, conflict, get_plan, mark_picked_up, settings_row
from app.services.workforce.lines import insert_line
from app.services.workforce.validation import parse_import_row, validate_line

KIND = "workforce_plan_lines"
MAX_ROWS = 2000
MAX_BYTES = 5 * 1024 * 1024
COLUMNS = (
    "role",
    "skills",
    "headcount",
    "type",
    "location",
    "start_month",
    "budget",
    "bill_rate",
    "unit",
    "notes",
)
HEADER_ALIASES = {
    "role": "role",
    "role_title": "role",
    "title": "role",
    "skills": "skills",
    "headcount": "headcount",
    "count": "headcount",
    "type": "type",
    "line_type": "type",
    "location": "location",
    "start_month": "start_month",
    "start month": "start_month",
    "month": "start_month",
    "budget": "budget",
    "bill_rate": "bill_rate",
    "bill rate": "bill_rate",
    "unit": "unit",
    "rate_unit": "unit",
    "notes": "notes",
}
TEMPLATE = (
    ",".join(COLUMNS)
    + "\nJava Developer,Java:4; Spring:3,3,new,Bengaluru,2027-04,4500000.00,,annual,Core platform team\n"
)


def _bad_file(message: str) -> ValidationFailed:
    return ValidationFailed(message, code="import_file_invalid")


def _normalise(header: list[str], values: list[Any]) -> dict[str, Any]:
    row: dict[str, Any] = {}
    for i, h in enumerate(header):
        key = HEADER_ALIASES.get(h)
        if key and i < len(values):
            v = values[i]
            row[key] = "" if v is None else (v if isinstance(v, str) else str(v))
    return row


def read_rows(content: bytes, filename: str | None) -> list[tuple[int, dict[str, Any]]]:
    """Data rows with their sheet / CSV row number (header = row 1). XLSX cells are read as text."""
    if len(content) > MAX_BYTES:
        raise _bad_file("The file is larger than 5 MB")
    name = (filename or "").lower()
    raw: list[list[Any]]
    if name.endswith(".xlsx") or content[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise _bad_file("The XLSX file cannot be read") from exc
        sheet = wb.worksheets[0]
        raw = [list(r) for r in sheet.iter_rows(values_only=True)]
        for r in raw:  # month cells typed as dates in Excel -> YYYY-MM
            for i, v in enumerate(r):
                if hasattr(v, "year") and hasattr(v, "month"):
                    r[i] = f"{v.year:04d}-{v.month:02d}"
                elif isinstance(v, float):
                    r[i] = format(v, "f").rstrip("0").rstrip(".") if v == int(v) else repr(v)
    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise _bad_file("The CSV file must be UTF-8") from exc
        raw = [list(r) for r in csv.reader(io.StringIO(text))]
    if not raw:
        raise _bad_file("The file is empty")
    header = [str(h or "").strip().lower() for h in raw[0]]
    if "role" not in {HEADER_ALIASES.get(h) for h in header}:
        raise ValidationFailed(
            "The file needs a role column",
            code="import_columns",
            details={"required": ["role", "headcount", "start_month"]},
        )
    out: list[tuple[int, dict[str, Any]]] = []
    for n, values in enumerate(raw[1:], start=2):
        if not any(str(v or "").strip() for v in values):
            continue
        out.append((n, _normalise(header, values)))
        if len(out) > MAX_ROWS:
            raise _bad_file(f"At most {MAX_ROWS} rows can be imported at once")
    return out


async def import_lines(
    ctx: RequestContext,
    db: AsyncSession,
    plan_id: uuid.UUID,
    content: bytes,
    filename: str | None,
    *,
    apply: bool = False,
) -> LineImportOut:
    from app.services import imports_service

    plan = await get_plan(ctx, db, plan_id, permission=WRITE, lock=True)
    if plan.status not in EDITABLE:
        raise conflict(
            "plan_not_editable",
            "Lines can only be imported while the plan is Draft, Budgeted or a Revision Draft",
            status=plan.status,
        )
    rows = read_rows(content, filename)
    settings = await settings_row(db)
    errors: list[ImportRowError] = []
    valid: list[dict[str, Any]] = []
    for n, row in rows:
        data, row_errors = parse_import_row(row)
        values, line_errors = await validate_line(
            ctx, db, plan, data, max_headcount=settings.max_headcount_per_line
        )
        problems = row_errors + [e for e in line_errors if not any(e[0] == r[0] for r in row_errors)]
        if problems:
            errors += [ImportRowError(row=n, field=f, message=m) for f, m in problems]
        else:
            valid.append(values)
    bad_rows = len({e.row for e in errors})
    digest = hashlib.sha256(plan.id.bytes + content).hexdigest()
    batch = await imports_service.register_preview(
        ctx,
        db,
        KIND,
        None,
        digest,
        len(rows),
        bad_rows,
        {"plan_id": str(plan.id), "errors": [e.model_dump() for e in errors[:500]], "valid": len(valid)},
    )
    created = 0
    if apply:
        await imports_service.mark_applied(ctx, db, batch.id)
        for values in valid:
            await insert_line(db, plan, values)
            created += 1
        mark_picked_up(ctx, plan)
        await db.flush()
        audit.record(
            "workforce.lines_imported",
            "workforce_plan",
            plan.id,
            session=db,
            batch_id=str(batch.id),
            created=created,
        )
    return LineImportOut(
        batch_id=batch.id,
        applied=apply,
        row_count=len(rows),
        valid_count=len(valid),
        error_count=bad_rows,
        created=created,
        errors=errors,
    )


async def import_plan_line(
    ctx: RequestContext, db: AsyncSession, data: PlanLineImportIn, *, validate_only: bool = False
) -> uuid.UUID | None:
    """Go-live import of one line into the editable plan named ``plan_name`` (FR-263)."""

    def bad(field: str, message: str) -> ValidationFailed:
        return ValidationFailed(
            message, code="import_invalid", details={"fields": [{"field": field, "message": message}]}
        )

    if not data.plan_name:
        raise bad("plan_name", "Name the plan the line belongs to")
    plans = list(
        (await db.execute(select(WorkforcePlan).where(WorkforcePlan.name == data.plan_name))).scalars()
    )
    if len(plans) != 1:
        raise bad("plan_name", "Unknown plan" if not plans else "Several plans have this name")
    plan = plans[0]
    if plan.status not in EDITABLE:
        raise bad("plan_name", "The plan is not editable")
    parsed, row_errors = parse_import_row(data.model_dump())
    settings = await settings_row(db)
    values, errors = await validate_line(ctx, db, plan, parsed, max_headcount=settings.max_headcount_per_line)
    problems = row_errors + errors
    if problems:
        field, message = problems[0]
        column = {
            "role_title": "role",
            "line_type": "type",
            "budget_amount": "budget",
            "bill_rate_amount": "bill_rate",
            "rate_unit": "unit",
            "location_text": "location",
        }.get(field.split(".")[0], field.split(".")[0])
        raise bad(column, message)
    if validate_only:
        return None
    row = await insert_line(db, plan, values)
    return row.id
