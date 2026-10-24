"""Payroll export files (Epic 28 A6): CSV / XLSX through a mapping template.

* Money / hours are written from ``Decimal`` only: CSV text decimal strings (``format(d, 'f')``), XLSX numeric
  cells holding the ``Decimal`` itself with ``number_format`` ``0.00`` (``0`` for zero-decimal currencies) --
  never through ``float``.
* Statutory placeholder columns are always empty strings (export only, PRD OQ2).
* PAN / bank are last-4 masks unless the requester holds ``identity:view``: then the full values are read via
  ``placements_service.worker_payment_identity(unmask=True)`` (audited by placements) and each download is
  audited ``payroll.file_unmasked``. Unmasked files are never cached; masked renders are cached as a
  ``stored_files`` row per version + mapping + format.
"""

import io
import uuid
from decimal import Decimal
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import audit, storage
from app.core.context import RequestContext
from app.core.money import minor_units
from app.models.files import StoredFile
from app.models.payroll import PayPeriod, PayrollExportMapping, PayrollExportRow, PayrollExportVersion
from app.services import files_service
from app.services.payroll import exports, mappings
from app.services.payroll.earnings import PLACEHOLDERS
from app.services.reporting.export import csv_line

FILE_OWNER = "payroll_export_version"
CONTENT_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
IDENTITY_VIEW = "identity:view"


def _mask(last4: str | None, width: int) -> str:
    return f"{'X' * width}{last4}" if last4 else ""


def field_value(
    key: str,
    row: PayrollExportRow,
    v: PayrollExportVersion,
    period: PayPeriod,
    name: str | None,
    full: Any = None,
) -> Any:
    if key in PLACEHOLDERS:
        return ""
    if key == "worker_name":
        return name or ""
    if key == "pan":
        return (full.pan if full is not None and full.pan else None) or _mask(row.pan_last4, 6)
    if key == "bank_account":
        return (full.account_number if full is not None and full.account_number else None) or _mask(
            row.bank_last4, 8
        )
    if key == "period_start":
        return period.period_start.isoformat()
    if key == "period_end":
        return period.period_end.isoformat()
    if key == "version_kind":
        return v.kind
    if key == "version_seq":
        return v.seq
    value = getattr(row, key, None)
    return "" if value is None else value


def _columns(m: PayrollExportMapping) -> list[tuple[str, str]]:
    return [(c["source_field"], c["header"]) for c in (m.columns or []) if c.get("include", True)]


def render_csv(header: list[str], rows: list[list[Any]]) -> bytes:
    out = io.StringIO()
    out.write(csv_line(header))
    for r in rows:
        out.write(csv_line(r))
    return out.getvalue().encode("utf-8")


def render_xlsx(header: list[str], rows: list[list[Any]], currencies: list[str]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r, ccy in zip(rows, currencies, strict=True):
        ws.append([v if isinstance(v, (Decimal, int)) else ("" if v is None else str(v)) for v in r])
        fmt = "0" if minor_units(ccy) == 0 else "0.00"
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, Decimal):
                cell.number_format = fmt
            elif isinstance(cell.value, str):
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def build_file(
    ctx: RequestContext,
    db: AsyncSession,
    version_id: uuid.UUID,
    mapping_id: uuid.UUID | None,
    fmt: str | None,
) -> files_service.Download:
    v = await exports.load_version(db, version_id)
    period = await db.get(PayPeriod, v.pay_period_id)
    assert period is not None
    m = await mappings.resolve(db, mapping_id, v.legal_entity_id)
    fmt = fmt or m.format
    principal = ctx.principal
    has = getattr(principal, "has", None)
    unmask = bool(callable(has) and has(IDENTITY_VIEW))
    filename = f"{v.human_id}-{m.id.hex[:8]}.{fmt}"
    if not unmask:
        cached = (
            await db.execute(
                select(StoredFile).where(
                    StoredFile.owner_entity_type == FILE_OWNER,
                    StoredFile.owner_entity_id == v.id,
                    StoredFile.original_name == filename,
                    StoredFile.deleted_at.is_(None),
                )
            )
        ).scalar_one_or_none()
        if cached is not None and v.status not in ("building",) and v.kind != "preview":
            audit.record(
                "payroll.file_downloaded", FILE_OWNER, v.id, session=db, mapping_id=str(m.id), format=fmt
            )
            return files_service.Download(
                data=await files_service.read_file(cached), filename=filename, content_type=CONTENT_TYPES[fmt]
            )
    rows = await exports.version_rows(db, v.id)
    names = await exports.worker_names(db, [r.candidate_id for r in rows])
    cols = _columns(m)
    header = [h for _, h in cols]
    data_rows: list[list[Any]] = []
    for r in rows:
        full = await exports._identity(ctx, db, r.candidate_id, unmask=True) if unmask else None
        values = [field_value(k, r, v, period, names.get(r.candidate_id), full) for k, _ in cols]
        data_rows.append(values)
    data = (
        render_csv(header, data_rows)
        if fmt == "csv"
        else render_xlsx(header, data_rows, [r.currency for r in rows])
    )
    if unmask:
        audit.record(
            "payroll.file_unmasked",
            FILE_OWNER,
            v.id,
            session=db,
            mapping_id=str(m.id),
            format=fmt,
            rows=len(rows),
        )
    else:
        audit.record(
            "payroll.file_downloaded", FILE_OWNER, v.id, session=db, mapping_id=str(m.id), format=fmt
        )
        if v.kind != "preview":
            await files_service.store_generated(
                ctx,
                db,
                data,
                key=storage.new_key(v.tenant_id),
                filename=filename,
                content_type=CONTENT_TYPES[fmt],
                owner_entity_type=FILE_OWNER,
                owner_entity_id=v.id,
            )
    return files_service.Download(data=data, filename=filename, content_type=CONTENT_TYPES[fmt])


async def file_access(ctx: RequestContext, db: AsyncSession, version_id: uuid.UUID, write: bool) -> bool:  # ruff: ignore[unused-function-argument]  # NOSONAR
    """``files_service`` checker ``payroll_export_version``: internal users with ``payroll:read``."""
    from app.core.permissions import Principal

    p = ctx.principal
    if p is None:
        return True
    if not isinstance(p, Principal) or write:
        return False
    return p.type == "user" and p.scope_of("payroll:read") is not None
