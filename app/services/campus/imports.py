# ruff: file-ignore[line-too-long]
"""Placement-officer list import (Epic 43 story 43.2, FR-218).

CSV / XLSX (``email``, ``name`` / ``full_name`` and optional standard answers ``program``, ``branch``, ``cgpa``,
``percentage``, ``backlogs``, ``graduation_year`` used as prefill) -> ``preview``: row errors (missing / invalid email,
invalid answers, duplicates in the file), recorded through ``imports_service.register_preview`` (kind
``campus_registration_invites``) with the valid rows kept as ``pending`` invites -> ``apply``: per row a personal
single-use Secure Link (purpose ``campus_registration``, subject ``registration_invite``) emailed to the student.

Imported rows never create Candidates or consents: that happens only when the student submits the form (consent is
never inferred). Invites expire with the registration window; their email / name are nulled when they expire
(``expire_invites``). The invitation email is sent directly (kind ``campus_registration_invite``) rather than through
``comms_service.queue_email`` so the raw link token is never stored in ``email_messages`` (deviation, recorded).
"""

import csv
import hashlib
import io
import uuid
from datetime import datetime, timedelta
from typing import Any

from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import audit, secure_links
from app.core.context import RequestContext
from app.core.errors import NotFound, ValidationFailed
from app.models.campus import CampusDrive, DriveInstitution, Institution, RegistrationInvite
from app.schemas.campus import ImportApplyOut, ImportPreviewOut, ImportRowOut
from app.services.campus import forms
from app.services.campus.common import (
    LINK_REGISTRATION,
    SUBJECT_INVITE,
    conflict,
    drive_for,
    now,
    registration_url,
    user_id_of,
)

KIND = "campus_registration_invites"
MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
ALIASES = {
    "email": "email",
    "email address": "email",
    "e-mail": "email",
    "name": "name",
    "full name": "name",
    "full_name": "name",
    "student name": "name",
    "program": "program",
    "programme": "program",
    "branch": "branch",
    "cgpa": "cgpa",
    "percentage": "percentage",
    "backlogs": "backlogs",
    "graduation_year": "graduation_year",
    "graduation year": "graduation_year",
}
PREFILL_KEYS = ("program", "branch", "cgpa", "percentage", "backlogs", "graduation_year")


def _bad(message: str) -> ValidationFailed:
    return ValidationFailed(message, code="import_file_invalid")


def read_rows(content: bytes, filename: str | None) -> list[tuple[int, dict[str, str]]]:
    if len(content) > MAX_BYTES:
        raise _bad("The file is larger than 5 MB")
    raw: list[list[Any]]
    if (filename or "").lower().endswith(".xlsx") or content[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise _bad("The XLSX file cannot be read") from exc
        raw = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise _bad("The CSV file must be UTF-8") from exc
        raw = [list(r) for r in csv.reader(io.StringIO(text))]
    if not raw:
        raise _bad("The file is empty")
    header = [ALIASES.get(str(h or "").strip().lower()) for h in raw[0]]
    if "email" not in header:
        raise ValidationFailed("The file needs an email column", code="import_columns", details={"required": ["email"]})
    out: list[tuple[int, dict[str, str]]] = []
    for n, values in enumerate(raw[1:], start=2):
        if not any(str(v or "").strip() for v in values):
            continue
        row: dict[str, str] = {}
        for i, key in enumerate(header):
            if key and i < len(values) and values[i] is not None:
                v = values[i]
                row[key] = (format(v, "f").rstrip("0").rstrip(".") if isinstance(v, float) else str(v)).strip()
        out.append((n, row))
        if len(out) > MAX_ROWS:
            raise _bad(f"At most {MAX_ROWS} rows can be imported at once")
    return out


def _prefill_errors(form: dict[str, Any], programs: list[dict[str, Any]], prefill: dict[str, str]) -> tuple[dict[str, Any], list[str]]:
    if not prefill:
        return {}, []
    try:
        clean, _ = forms.validate_answers(form, prefill, programs, partial=True)
    except ValidationFailed as exc:
        fields = (exc.details or {}).get("fields") or {}
        return {}, [f"{k}: {v}" for k, v in fields.items()]
    return clean, []


async def preview(
    c: RequestContext, db: AsyncSession, drive_id: uuid.UUID, drive_institution_id: uuid.UUID, content: bytes, filename: str | None
) -> ImportPreviewOut:
    from app.services import imports_service
    from app.services.candidates.common import normalize_email

    drive = await drive_for(c, db, drive_id, write=True)
    if drive.status not in ("draft", "registration_open", "in_progress"):
        raise conflict("drive_closed", "This drive is closed", status=drive.status)
    di = await db.get(DriveInstitution, drive_institution_id)
    if di is None or di.drive_id != drive.id:
        raise NotFound("Institution not found on this drive")
    inst = await db.get(Institution, di.institution_id)
    programs = list(inst.programs or []) if inst is not None else []
    rows = read_rows(content, filename)
    seen: dict[str, int] = {}
    out_rows: list[ImportRowOut] = []
    valid: list[tuple[str, str | None, dict[str, Any]]] = []
    for n, row in rows:
        errors: list[str] = []
        raw_email = row.get("email") or ""
        email = normalize_email(raw_email) if raw_email else None
        if not raw_email:
            errors.append("email: missing")
        elif email is None:
            errors.append("email: invalid")
        elif email in seen:
            errors.append(f"email: duplicate of row {seen[email]}")
        if email and email not in seen:
            seen[email] = n
        prefill, perrs = _prefill_errors(drive.form_schema, programs, {k: row[k] for k in PREFILL_KEYS if row.get(k)})
        errors += perrs
        name = (row.get("name") or "").strip()[:200] or None
        out_rows.append(ImportRowOut(row=n, email=email or raw_email or None, name=name, errors=errors))
        if not errors and email:
            valid.append((email, name, prefill))
    batch = await imports_service.register_preview(
        c, db, KIND, None, hashlib.sha256(content).hexdigest(), len(rows), sum(1 for r in out_rows if r.errors),
        {"drive_id": str(drive.id), "drive_institution_id": str(di.id), "valid": len(valid)},
    )
    expires = drive.registration_closes_at or (now() + timedelta(days=30))
    for email, name, prefill in valid:
        db.add(
            RegistrationInvite(
                id=uuid.uuid4(), tenant_id=drive.tenant_id, drive_id=drive.id, drive_institution_id=di.id, email=email,
                name=name, prefill={k: v for k, v in prefill.items() if k != "__file__"}, import_batch_id=batch.id,
                status="pending", expires_at=expires,
            )
        )
    await db.flush()
    return ImportPreviewOut(
        batch_id=batch.id, drive_institution_id=di.id, row_count=len(rows),
        error_count=sum(1 for r in out_rows if r.errors), rows=out_rows,
    )


async def _send_invite(db: AsyncSession, drive: CampusDrive, inst_name: str, invite: RegistrationInvite, token: str) -> None:
    from app.core.mailer import EmailMessage, send_email
    from app.services import identity_bridge

    tenant = await identity_bridge.get_tenant(db, drive.tenant_id)
    brand = tenant.name if tenant is not None else "our team"
    closes = f" before {drive.registration_closes_at:%d %b %Y}" if drive.registration_closes_at else ""
    await send_email(
        EmailMessage(
            to=str(invite.email),
            subject=f"Register for the {brand} campus drive: {drive.name}"[:200],
            text=(
                f"Hello {(invite.name or '').split(' ')[0] or 'there'},\n\n"
                f"Your placement office at {inst_name} shared your details so that we could invite you to "
                f"{drive.name} (batch {drive.batch_year}). We have not created a profile for you.\n\n"
                f"If you want to take part, register here{closes}:\n{registration_url(token)}\n\n"
                "The registration page explains how we use your data. If you do not register, we keep nothing "
                "about you beyond this invitation, which expires with the registration window.\n\n"
                f"{brand}\n"
            ),
            tenant_id=drive.tenant_id,
            kind="campus_registration_invite",
            from_name=brand,
        )
    )


async def apply(c: RequestContext, db: AsyncSession, drive_id: uuid.UUID, batch_id: uuid.UUID) -> ImportApplyOut:
    from app.services import imports_service

    drive = await drive_for(c, db, drive_id, write=True)
    if drive.status not in ("registration_open", "in_progress"):
        raise conflict("drive_not_open", "Open registration before sending invitations", status=drive.status)
    if drive.registration_closes_at is not None and drive.registration_closes_at <= now():
        raise conflict("registration_closed", "Registration for this drive is closed")
    invites = list(
        (
            await db.execute(
                select(RegistrationInvite)
                .where(RegistrationInvite.import_batch_id == batch_id, RegistrationInvite.drive_id == drive.id)
                .with_for_update()
            )
        ).scalars()
    )
    if not invites:
        raise NotFound("Import batch not found")
    await imports_service.mark_applied(c, db, batch_id)
    di_ids = {i.drive_institution_id for i in invites}
    inst_names = {
        d.id: str(n)
        for d, n in (
            await db.execute(
                select(DriveInstitution, Institution.name)
                .join(Institution, Institution.id == DriveInstitution.institution_id)
                .where(DriveInstitution.id.in_(di_ids))
            )
        ).tuples()
    }
    from app.models.campus import Registration
    from app.models.candidates import CandidateEmail

    already = {
        str(e).lower()
        for e in (
            await db.execute(
                select(CandidateEmail.email)
                .join(Registration, Registration.candidate_id == CandidateEmail.candidate_id)
                .where(Registration.drive_id == drive.id, CandidateEmail.email.in_([i.email for i in invites if i.email]))
            )
        ).scalars()
    }
    sent = skipped = 0
    for inv in invites:
        if inv.status != "pending" or inv.email is None or str(inv.email).lower() in already:
            skipped += 1
            continue
        ttl = inv.expires_at + timedelta(days=1) - now()
        link, token = await secure_links.issue(
            db, tenant_id=drive.tenant_id, purpose=LINK_REGISTRATION, subject_type=SUBJECT_INVITE, subject_id=inv.id,
            ttl=ttl if ttl > timedelta(hours=1) else timedelta(hours=1), max_uses=1,
            metadata={"drive_id": str(drive.id)}, created_by=user_id_of(c),
        )
        inv.secure_link_id = link.id
        inv.status = "sent"
        await _send_invite(db, drive, inst_names.get(inv.drive_institution_id, "your institution"), inv, token)
        sent += 1
    await db.flush()
    audit.record("campus.invites_sent", "campus_drive", drive.id, session=db, sent=sent, skipped=skipped)
    return ImportApplyOut(batch_id=batch_id, sent=sent, skipped=skipped)


async def revoke_invites(db: AsyncSession, drive_id: uuid.UUID) -> int:
    rows = list(
        (
            await db.execute(
                select(RegistrationInvite).where(
                    RegistrationInvite.drive_id == drive_id, RegistrationInvite.status.in_(("pending", "sent"))
                )
            )
        ).scalars()
    )
    for inv in rows:
        if inv.secure_link_id is not None:
            await secure_links.revoke(db, inv.secure_link_id)
        inv.status = "expired"
        inv.email = None
        inv.name = None
        inv.prefill = {}
    await db.flush()
    return len(rows)


async def expire_invites(db: AsyncSession, at: datetime | None = None) -> int:
    """Minimise invites whose window ended (pending previews older than a day are dropped as well)."""
    moment = at or now()
    res = await db.execute(
        update(RegistrationInvite)
        .where(
            RegistrationInvite.status.in_(("pending", "sent")),
            (RegistrationInvite.expires_at <= moment)
            | ((RegistrationInvite.status == "pending") & (RegistrationInvite.created_at <= moment - timedelta(days=1))),
        )
        .values(status="expired", email=None, name=None, prefill={})
    )
    return int(getattr(res, "rowcount", 0) or 0)
