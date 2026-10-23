# ruff: file-ignore[line-too-long, unused-function-argument]
"""Audiences (Epic 39 A2, FR-196, story 39.1).

``resolve`` = union of the include sources (Pools via ``pools_service.member_ids``, Segments via
``segments_service.evaluate`` -- never compiled here) and the include list, minus the same for exclude; always under
the caller's data scope (the Epic 38 services apply it; list entries are re-filtered here), erased / merged candidates
dropped. ``preview`` runs ``governance.bulk_preview`` (the same rules, set-based). ``freeze`` stores the resolved ids
as a snapshot (``audience_members``, batches of 5,000; AD-43). Consent is evaluated at send time, never as an
audience filter (AD-42).
"""

import csv
import io
import time
import uuid
from collections.abc import Iterable
from typing import Any

from sqlalchemy import delete, func, insert, select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.context import RequestContext
from app.core.errors import NotFound, ValidationFailed
from app.core.permissions import Principal, scope_condition
from app.models.campaigns import (
    CHANNELS,
    Audience,
    AudienceListEntry,
    AudienceMember,
    AudienceSnapshot,
    AudienceSource,
)
from app.models.candidates import Candidate, CandidateEmail, CandidatePhone
from app.models.talent_pools import Segment, TalentPool
from app.schemas.campaigns import (
    AudienceCreate,
    AudienceOut,
    AudiencePreviewOut,
    AudienceSourceIn,
    AudienceSourceOut,
    AudienceUpdate,
    ListEntriesIn,
    ListImportOut,
    ListImportRow,
)
from app.services.campaigns import governance
from app.services.campaigns.common import not_found, now, principal_of, tenant_of, user_id_of

FREEZE_BATCH = 5000
MAX_IMPORT_ROWS = 50_000


# ================================================================ scope
def _candidate_scope(ctx: RequestContext) -> Any:
    base = (Candidate.merged_into_id.is_(None)) & (Candidate.erased_at.is_(None))
    p = ctx.principal
    if not isinstance(p, Principal):  # system context (tasks): RLS still applies
        return base
    cond = scope_condition(p, "candidate:read", Candidate.owner_id)
    return base if cond is None else base & cond


async def _scoped(ctx: RequestContext, db: AsyncSession, ids: Iterable[uuid.UUID]) -> set[uuid.UUID]:
    wanted = list(set(ids))
    out: set[uuid.UUID] = set()
    for i in range(0, len(wanted), 20_000):
        part = wanted[i : i + 20_000]
        out |= set(
            (
                await db.execute(select(Candidate.id).where(Candidate.id.in_(part), _candidate_scope(ctx)))
            ).scalars()
        )
    return out


# ================================================================ CRUD
async def _get(db: AsyncSession, audience_id: uuid.UUID) -> Audience:
    row = await db.get(Audience, audience_id)
    if row is None:
        raise not_found("Audience")
    return row


async def _sources(db: AsyncSession, audience_id: uuid.UUID) -> list[AudienceSource]:
    return list(
        (
            await db.execute(
                select(AudienceSource)
                .where(AudienceSource.audience_id == audience_id)
                .order_by(AudienceSource.created_at)
            )
        ).scalars()
    )


async def to_out(db: AsyncSession, row: Audience) -> AudienceOut:
    await db.flush()
    await db.refresh(row)
    sources = await _sources(db, row.id)
    pool_ids = {s.talent_pool_id for s in sources if s.talent_pool_id}
    seg_ids = {s.segment_id for s in sources if s.segment_id}
    pools = (
        dict(
            (await db.execute(select(TalentPool.id, TalentPool.name).where(TalentPool.id.in_(pool_ids))))
            .tuples()
            .all()
        )
        if pool_ids
        else {}
    )
    segs = (
        dict(
            (await db.execute(select(Segment.id, Segment.name).where(Segment.id.in_(seg_ids)))).tuples().all()
        )
        if seg_ids
        else {}
    )
    counts = dict(
        (
            await db.execute(
                select(AudienceListEntry.mode, func.count())
                .where(AudienceListEntry.audience_id == row.id)
                .group_by(AudienceListEntry.mode)
            )
        )
        .tuples()
        .all()
    )
    return AudienceOut(
        id=row.id,
        name=row.name,
        description=row.description,
        owner_id=row.owner_id,
        status=row.status,
        version=row.version,
        sources=[
            AudienceSourceOut(
                id=s.id,
                mode=s.mode,
                kind=s.kind,
                talent_pool_id=s.talent_pool_id,
                segment_id=s.segment_id,
                name=pools.get(s.talent_pool_id) if s.kind == "pool" else segs.get(s.segment_id),
            )
            for s in sources
        ],
        include_list_count=int(counts.get("include", 0)),
        exclude_list_count=int(counts.get("exclude", 0)),
        created_at=row.created_at,
        updated_at=row.updated_at,
    )


async def _set_sources(
    ctx: RequestContext, db: AsyncSession, audience: Audience, sources: list[AudienceSourceIn]
) -> None:
    from app.services import pools_service

    for i, s in enumerate(sources):
        ref = s.talent_pool_id if s.kind == "pool" else s.segment_id
        other = s.segment_id if s.kind == "pool" else s.talent_pool_id
        if ref is None or other is not None:
            raise ValidationFailed(
                "Each source names exactly one pool or segment",
                code="invalid_audience_source",
                details={"fields": {f"sources[{i}]": "Exactly one reference matching kind"}},
            )
        if s.kind == "pool":
            await pools_service.pool_summary(ctx, db, ref)  # 404 outside the viewer's pool visibility
        elif (await db.execute(select(Segment.id).where(Segment.id == ref))).first() is None:
            raise NotFound("Segment not found")
    await db.execute(delete(AudienceSource).where(AudienceSource.audience_id == audience.id))
    for s in sources:
        db.add(
            AudienceSource(
                id=uuid.uuid4(),
                tenant_id=audience.tenant_id,
                audience_id=audience.id,
                mode=s.mode,
                kind=s.kind,
                talent_pool_id=s.talent_pool_id if s.kind == "pool" else None,
                segment_id=s.segment_id if s.kind == "segment" else None,
            )
        )
    await db.flush()


async def list_audiences(
    ctx: RequestContext, db: AsyncSession, include_archived: bool = False
) -> list[AudienceOut]:
    stmt = select(Audience).order_by(Audience.updated_at.desc()).limit(500)
    if not include_archived:
        stmt = stmt.where(Audience.status != "archived")
    return [await to_out(db, r) for r in (await db.execute(stmt)).scalars()]


async def create_audience(ctx: RequestContext, db: AsyncSession, data: AudienceCreate) -> AudienceOut:
    principal = principal_of(ctx)
    row = Audience(
        id=uuid.uuid4(),
        tenant_id=tenant_of(db),
        name=data.name,
        description=data.description,
        owner_id=principal.id,
        status="active",
        created_by=user_id_of(ctx),
    )
    db.add(row)
    await db.flush()
    await _set_sources(ctx, db, row, data.sources)
    await db.refresh(row)
    return await to_out(db, row)


async def get_audience(ctx: RequestContext, db: AsyncSession, audience_id: uuid.UUID) -> AudienceOut:
    return await to_out(db, await _get(db, audience_id))


async def update_audience(
    ctx: RequestContext, db: AsyncSession, audience_id: uuid.UUID, data: AudienceUpdate
) -> AudienceOut:
    row = await _get(db, audience_id)
    if row.status == "archived":
        from app.services.campaigns.common import conflict

        raise conflict("audience_archived", "The audience is archived")
    if data.name is not None:
        row.name = data.name
    if data.description is not None:
        row.description = data.description or None
    if data.status is not None:
        row.status = data.status
    if data.sources is not None:
        await _set_sources(ctx, db, row, data.sources)
    row.version += 1
    row.updated_at = now()
    await db.flush()
    await db.refresh(row)
    return await to_out(db, row)


async def archive_audience(ctx: RequestContext, db: AsyncSession, audience_id: uuid.UUID) -> None:
    row = await _get(db, audience_id)
    row.status = "archived"
    await db.flush()


# ================================================================ list entries
async def add_list_entries(
    ctx: RequestContext, db: AsyncSession, audience_id: uuid.UUID, data: ListEntriesIn
) -> AudienceOut:
    row = await _get(db, audience_id)
    allowed = await _scoped(ctx, db, data.candidate_ids)
    missing = [i for i in data.candidate_ids if i not in allowed]
    if missing:
        raise NotFound("Candidate not found", details={"candidate_id": str(missing[0])})
    await _insert_entries(db, row, data.mode, sorted(allowed))
    return await to_out(db, row)


async def remove_list_entry(
    ctx: RequestContext, db: AsyncSession, audience_id: uuid.UUID, mode: str, candidate_id: uuid.UUID
) -> None:
    await _get(db, audience_id)
    await db.execute(
        delete(AudienceListEntry).where(
            AudienceListEntry.audience_id == audience_id,
            AudienceListEntry.mode == mode,
            AudienceListEntry.candidate_id == candidate_id,
        )
    )


async def _insert_entries(db: AsyncSession, audience: Audience, mode: str, ids: list[uuid.UUID]) -> None:
    for i in range(0, len(ids), FREEZE_BATCH):
        part = ids[i : i + FREEZE_BATCH]
        await db.execute(
            pg_insert(AudienceListEntry)
            .values(
                [
                    {
                        "tenant_id": audience.tenant_id,
                        "audience_id": audience.id,
                        "mode": mode,
                        "candidate_id": c,
                    }
                    for c in part
                ]
            )
            .on_conflict_do_nothing()
        )
    await db.flush()


def _rows_from_upload(filename: str | None, content: bytes) -> list[str]:
    name = (filename or "").lower()
    values: list[str] = []
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for row in wb.active.iter_rows(values_only=True):
            cell = next((c for c in row if c not in (None, "")), None)
            values.append("" if cell is None else str(cell).strip())
    else:
        text = content.decode("utf-8-sig", errors="replace")
        for row in csv.reader(io.StringIO(text)):
            cell = next((c for c in row if c.strip()), "")
            values.append((cell or "").strip())
    if values and values[0].lower() in ("email", "phone", "email or phone", "value", "contact"):
        values[0] = ""
    return values[:MAX_IMPORT_ROWS]


async def list_import(
    ctx: RequestContext,
    db: AsyncSession,
    audience_id: uuid.UUID,
    filename: str | None,
    content: bytes,
    *,
    mode: str,
    apply: bool,
) -> ListImportOut:
    """Match uploaded emails / E.164 phones to Candidates (never creates one); ``apply`` commits the matched rows."""
    from app.services.candidates.common import normalize_phone

    audience = await _get(db, audience_id)
    values = _rows_from_upload(filename, content)
    emails: dict[int, str] = {}
    phones: dict[int, str] = {}
    for i, v in enumerate(values, start=1):
        if not v:
            continue
        if "@" in v:
            emails[i] = v.lower()
        else:
            p = normalize_phone(v)
            if p:
                phones[i] = p
    by_email: dict[str, uuid.UUID] = {}
    by_phone: dict[str, uuid.UUID] = {}
    scope = _candidate_scope(ctx)
    if emails:
        rows = await db.execute(
            select(func.lower(CandidateEmail.email), CandidateEmail.candidate_id)
            .join(Candidate, Candidate.id == CandidateEmail.candidate_id)
            .where(func.lower(CandidateEmail.email).in_(sorted(set(emails.values()))), scope)
        )
        by_email = {str(e): cid for e, cid in rows.all()}
    if phones:
        rows = await db.execute(
            select(CandidatePhone.phone_e164, CandidatePhone.candidate_id)
            .join(Candidate, Candidate.id == CandidatePhone.candidate_id)
            .where(CandidatePhone.phone_e164.in_(sorted(set(phones.values()))), scope)
        )
        by_phone = {str(p): cid for p, cid in rows.all()}
    matched: list[ListImportRow] = []
    unmatched: list[ListImportRow] = []
    for i, v in enumerate(values, start=1):
        if not v:
            continue
        cid = by_email.get(emails[i]) if i in emails else by_phone.get(phones[i]) if i in phones else None
        (matched if cid else unmatched).append(ListImportRow(row=i, value=_mask(v), candidate_id=cid))
    if apply and matched:
        await _insert_entries(db, audience, mode, sorted({m.candidate_id for m in matched if m.candidate_id}))
    return ListImportOut(
        mode=mode,
        applied=apply,
        matched=len(matched),
        unmatched=len(unmatched),
        matched_rows=matched[:500],
        unmatched_rows=unmatched[:500],
    )


def _mask(value: str) -> str:
    if "@" in value:
        local, _, domain = value.partition("@")
        return f"{local[:2]}***@{domain}"
    return f"{value[:3]}***{value[-2:]}" if len(value) > 5 else "***"


# ================================================================ resolve / preview / freeze
async def _source_ids(ctx: RequestContext, db: AsyncSession, source: AudienceSource) -> set[uuid.UUID]:
    from app.services import pools_service, segments_service

    if source.kind == "pool" and source.talent_pool_id is not None:
        out: set[uuid.UUID] = set()
        after: uuid.UUID | None = None
        while True:
            page = await pools_service.member_ids(ctx, db, source.talent_pool_id, limit=20_000, after=after)
            out.update(page)
            if len(page) < 20_000:
                return out
            after = page[-1]
    if source.kind == "segment" and source.segment_id is not None:
        evaluation = await segments_service.evaluate(ctx, db, source.segment_id, mode="ids")
        return set(evaluation.subject_ids)
    return set()


async def resolve(ctx: RequestContext, db: AsyncSession, audience_id: uuid.UUID) -> set[uuid.UUID]:
    audience = await _get(db, audience_id)
    include: set[uuid.UUID] = set()
    exclude: set[uuid.UUID] = set()
    for source in await _sources(db, audience.id):
        (include if source.mode == "include" else exclude).update(await _source_ids(ctx, db, source))
    entries = await db.execute(
        select(AudienceListEntry.mode, AudienceListEntry.candidate_id).where(
            AudienceListEntry.audience_id == audience.id
        )
    )
    for mode, cid in entries.all():
        (include if mode == "include" else exclude).add(cid)
    wanted = include - exclude
    return await _scoped(ctx, db, wanted)


async def preview(
    ctx: RequestContext, db: AsyncSession, audience_id: uuid.UUID, channel: str | None = None
) -> AudiencePreviewOut:
    started = time.monotonic()
    ids = await resolve(ctx, db, audience_id)
    channels = (channel,) if channel else CHANNELS
    reachable, suppressed = await governance.bulk_preview(db, ids, channels, now())
    return AudiencePreviewOut(
        audience_id=audience_id,
        total=len(ids),
        reachable=reachable,
        suppressed=suppressed,
        elapsed_ms=int((time.monotonic() - started) * 1000),
    )


async def freeze(
    ctx: RequestContext,
    db: AsyncSession,
    audience_id: uuid.UUID,
    purpose: str,
    counts: dict[str, Any] | None = None,
) -> uuid.UUID:
    audience = await _get(db, audience_id)
    ids = sorted(await resolve(ctx, db, audience_id))
    snapshot = AudienceSnapshot(
        id=uuid.uuid4(),
        tenant_id=audience.tenant_id,
        audience_id=audience.id,
        purpose=purpose,
        taken_at=now(),
        total=len(ids),
        counts=counts or {},
    )
    db.add(snapshot)
    await db.flush()
    for i in range(0, len(ids), FREEZE_BATCH):
        part = ids[i : i + FREEZE_BATCH]
        await db.execute(
            insert(AudienceMember),
            [{"tenant_id": audience.tenant_id, "snapshot_id": snapshot.id, "candidate_id": c} for c in part],
        )
    await db.flush()
    return snapshot.id


async def snapshot_member_ids(db: AsyncSession, snapshot_id: uuid.UUID) -> list[uuid.UUID]:
    return list(
        (
            await db.execute(
                select(AudienceMember.candidate_id)
                .where(AudienceMember.snapshot_id == snapshot_id)
                .order_by(AudienceMember.candidate_id)
            )
        ).scalars()
    )
